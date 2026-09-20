import json

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from lxml import etree

from transmux.alignment import source_blocks, normalize_groups, window, group_mappings, target_texts
from transmux.documents import extract, export_groups, revise_group_docx, doc_paragraphs
from transmux.jobs import Worker
from transmux.store import Store


def fixture_doc(tmp_path):
    path = tmp_path / 'source.docx'
    doc = Document()
    doc.add_heading('标题', 1)
    doc.add_paragraph('模型通过分析')
    doc.add_paragraph('数据进行预测。')
    doc.add_paragraph('方法有效。还需要讨论其局限。')
    doc.add_paragraph('列表条目', style='List Bullet')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格数据'
    doc.save(path)
    return path, source_blocks(path, extract(path))


def group(ids, texts, reason=''):
    return dict(source_ids=ids, paragraphs=texts, reason=reason)


def test_merge_split_export_preserves_structure_and_revision(tmp_path):
    source, blocks = fixture_doc(tmp_path)
    original = source.read_bytes()
    values = [group(['p000001'], ['Title']), group(['p000002', 'p000003'], ['The model predicts by analyzing data.'], 'Repair a broken sentence.'),
              group(['p000004'], ['The method works.', 'Its limitations need discussion.'], 'Separate topics.'),
              group(['p000005'], ['List item']), group(['p000006'], ['Table data'])]
    rows = normalize_groups(values, blocks)
    output = tmp_path / 'translated.docx'
    export_groups(source, blocks, rows, output)
    assert extract(output) == target_texts(rows)
    doc = Document(output)
    assert doc.paragraphs[0].style.name == 'Heading 1'
    assert doc.paragraphs[-1].style.name == 'List Bullet'
    assert doc.tables[0].cell(0, 0).text == 'Table data'
    before = [etree.tostring(p._p) for p in doc_paragraphs(doc)]
    revised = tmp_path / 'revised.docx'
    revise_group_docx(output, rows, 1, ['The model analyzes data.', 'It makes predictions.'], revised)
    after = [etree.tostring(p._p) for p in doc_paragraphs(Document(revised))]
    assert before[0] == after[0] and before[2:] == after[3:]
    assert source.read_bytes() == original


@pytest.mark.parametrize('values', [
    [group(['p000001'], ['Title'])],
    [group(['p000002', 'p000001'], ['Title'], 'Reorder')],
    [group(['p000001', 'p000002'], ['Title'], 'Cross heading')],
    [group(['p000001'], ['First', 'Second'], 'Split heading')],
    [group(['p000001'], ['Title']), group(['p000001'], ['Repeated'])],
    [group(['p000001'], ['Title']), group(['p000002', 'p000003'], ['Merged'])],
    [group(['p000001'], ['Title']), group(['p000002'], ['A\nB'])],
])
def test_invalid_coverage_or_structure_rejected(tmp_path, values):
    _, blocks = fixture_doc(tmp_path)
    with pytest.raises(ValueError):
        normalize_groups(values, blocks)


def test_blank_objects_sections_and_literal_lists_protected(tmp_path):
    source = tmp_path / 'source.docx'
    doc = Document()
    doc.add_paragraph('First')
    doc.add_paragraph('')
    doc.add_paragraph('Second')
    doc.add_paragraph('New page').add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph('Caption', style='Caption')
    doc.add_paragraph('Item', style='List Number')
    doc.add_paragraph('1. Literal numbered item')
    doc.add_paragraph('Fig. 1 Caption without a caption style')
    doc.save(source)
    blocks = source_blocks(source, extract(source))
    assert blocks[0]['segment'] != blocks[1]['segment']
    assert all(b['protected'] for b in blocks[2:])


def test_window_carries_entire_trailing_group_and_no_eight_paragraph_limit(tmp_path):
    path = tmp_path / 'source.txt'
    texts = ['Text.'] * 12
    path.write_text('\n'.join(texts))
    blocks = source_blocks(path, texts)
    assert len(window(blocks, 0)[0]) == 12
    batch, continues = window(blocks, 0, max_chars=30)
    assert len(batch) == 6 and continues
    groups = normalize_groups([group([b['id']], [b['text']]) for b in batch[:4]] +
                              [group([b['id'] for b in batch[4:]], ['Merged.'], 'One idea')], batch, continues)
    offset = sum(len(g['source_ids']) for g in groups[:-1])
    assert offset == 4
    assert window(blocks, offset, 30)[0][0]['id'] == 'p000005'
    with pytest.raises(ValueError, match='at least two'):
        normalize_groups([group([b['id'] for b in batch], ['All'], 'One idea')], batch, True)


def test_mapping_evidence_uses_ids_inside_merged_group(tmp_path):
    source, blocks = fixture_doc(tmp_path)
    groups = normalize_groups([group(['p000002', 'p000003'], ['The model predicts by analyzing data.'], 'Repair')], blocks[1:3])
    candidates = [dict(original='数据', translation='data', context='Model input', source_id='p000003', target_id='p000002:t1'),
                  dict(original='数据', translation='data', context='Model input', source_id='p000002', target_id='p000002:t1')]
    accepted, report = group_mappings(Worker, candidates, groups, 'en')
    assert len(accepted) == len(report) == 1
    assert accepted[0]['source_id'] == 'p000003'


async def test_bad_structure_retries_and_group_revision_restores_original_count(tmp_path):
    store = Store(tmp_path / 'data')
    pid = store.create_project('groups', 'codex')['id']
    source = store.workspace(pid) / 'sources' / 'source.txt'
    source.write_text('模型通过分析\n数据进行预测。')
    source.with_suffix('.txt.json').write_text(json.dumps(extract(source)))
    fid = store.add_file(pid, source.name, 'source', source)

    class Runner:
        count = 0

        async def run(self, pid, jid, prompt, schema):
            props = schema['properties']
            if 'translations' in props:
                self.count += 1
                return dict(translations=[group(['p000002'] if self.count == 1 else ['p000001', 'p000002'],
                                               ['The model predicts by analyzing data.'], 'Repair')], mappings=[])
            if 'translation' in props:
                return dict(translation=['The model analyzes data.', 'It makes predictions.'])
            return dict(passed=True, issues=[], mapping_issues=[])

    worker = Worker(store, Runner())
    task = store.enqueue(pid, 'translate', dict(file_id=fid, use_rag=False, max_review_rounds=2))
    await worker.execute(task)
    result = store.rows('SELECT * FROM jobs WHERE id=?', (task['id'],))[0]
    assert result['state'] == 'succeeded', result['result']
    output = json.loads(result['result'])['file_id']
    assert len(store.comparison(pid, output)['paragraphs']) == 1
    task = store.enqueue(pid, 'revise', dict(file_id=output, paragraph=1, message='Restore original paragraph count.', use_rag=False, max_review_rounds=1))
    await worker.execute(task)
    result = store.rows('SELECT * FROM jobs WHERE id=?', (task['id'],))[0]
    assert result['state'] == 'succeeded', result['result']
    revised = json.loads(result['result'])['file_id']
    rows = store.comparison(pid, revised)['paragraphs']
    assert len(rows[0]['translations']) == 2
    assert rows[0]['source_ids'] == ['p000001', 'p000002']
    assert extract(store.download_path(pid, store.file(pid, revised))) == target_texts(rows)
    store.db.close()


async def test_cross_window_merge_has_no_duplicate_coverage_or_stale_mappings(tmp_path, monkeypatch):
    store = Store(tmp_path / 'data')
    pid = store.create_project('windows', 'codex')['id']
    source = store.workspace(pid) / 'sources' / 'source.txt'
    source.write_text('First.\nOther.\nModel.\nWorks.\nFinal.')
    source.with_suffix('.txt.json').write_text(json.dumps(extract(source)))
    fid = store.add_file(pid, source.name, 'source', source)
    from transmux.sections import ChapterPlanner
    monkeypatch.setattr('transmux.jobs.ChapterPlanner', lambda blocks: ChapterPlanner(blocks, max_chars=18))

    class Runner:
        batches = []

        async def run(self, pid, jid, prompt, schema):
            if 'translations' not in schema['properties']:
                return dict(passed=True, issues=[], mapping_issues=[])
            inputs = sorted((store.workspace(pid) / 'runs' / jid).glob('batch-*-input.json'))
            data = json.loads(inputs[-1].read_text())
            ids = [r['id'] for r in data['source_blocks']]
            self.batches.append(ids)
            if ids[0] == 'p000001':
                return dict(translations=[group([sid], [text]) for sid, text in zip(ids, data['source'])],
                            mappings=[dict(original='Model', translation='Model', context='Input',
                                           source_id='p000003', target_id='p000003:t1')])
            return dict(translations=[group(ids[:2], ['The model works.'], 'Sentence repair'),
                                      group(ids[2:], ['Final.'])], mappings=[])

    runner = Runner()
    worker = Worker(store, runner)
    task = store.enqueue(pid, 'translate', dict(file_id=fid, use_rag=False, max_review_rounds=1))
    await worker.execute(task)
    result = store.rows('SELECT * FROM jobs WHERE id=?', (task['id'],))[0]
    assert result['state'] == 'succeeded', result['result']
    rows = store.comparison(pid, json.loads(result['result'])['file_id'])['paragraphs']
    assert [sid for row in rows for sid in row['source_ids']] == [f'p{i:06d}' for i in range(1, 6)]
    assert rows[2]['source_ids'] == ['p000003', 'p000004']
    assert runner.batches == [['p000001', 'p000002', 'p000003'], ['p000003', 'p000004', 'p000005']]
    assert json.loads(result['result'])['new_mappings'] == 0
    assert target_texts(rows) == ['First.', 'Other.', 'The model works.', 'Final.']
    store.db.close()
