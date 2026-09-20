"""Manual browser integration: python tests/browser_smoke.py (requires playwright + Chrome)."""
import io
import asyncio
import tempfile
import threading
import time
from unittest.mock import patch

import httpx
import uvicorn
from docx import Document
from playwright.sync_api import sync_playwright

from transmux.app import create_app
from transmux.jobs import Worker
from transmux.rag import Rag
from test_workflows import FakeRunner, TinyEmbeddings
from test_layout import fake_pdf


class VisibleRunner(FakeRunner):
    async def run(self, *args, **kwargs):
        await asyncio.sleep(1.5)
        schema = kwargs.get('schema') or (args[3] if len(args) > 3 else None)
        if schema and 'translation' in schema['properties']:
            return {'translation':'The physician works at the hospital.'}
        return await super().run(*args, **kwargs)


def main():
    with tempfile.TemporaryDirectory(prefix="transmux-browser-") as root, patch('transmux.layout.render_pdf', fake_pdf):
        app = create_app(root, lambda store: Worker(store, VisibleRunner(), Rag(TinyEmbeddings())))
        server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=18765, log_level="error"))
        thread = threading.Thread(target=server.run)
        thread.start()
        try:
            for _ in range(100):
                try:
                    if httpx.get('http://127.0.0.1:18765/api/health').status_code == 200:
                        break
                except httpx.ConnectError:
                    pass
                time.sleep(.1)
            with sync_playwright() as p:
                browser = p.chromium.launch(executable_path="/usr/bin/google-chrome", headless=True, args=['--no-sandbox'])
                page = browser.new_page(viewport={"width":1440,"height":1050}, device_scale_factor=1)
                errors = []
                page.on('pageerror', lambda e: errors.append(str(e)))
                page.goto('http://127.0.0.1:18765')
                page.locator('#start-project').click()
                page.locator('#project-name').fill('产品文档 · 中英翻译')
                page.wait_for_function("document.querySelector('#create-model-options').options.length > 2")
                page.locator('#create-model-options').select_option('__custom__')
                page.locator('#create-model').fill('test-initial-model')
                page.locator('#project-form button[type=submit]').click()
                page.locator('#workspace').wait_for(state='visible')
                assert page.locator('#project-model').input_value() == 'test-initial-model'
                assert page.locator('[data-scroll]').evaluate_all('els => els.map(el => el.dataset.scroll)') == ['corpus', 'translation', 'knowledge', 'library', 'conversation']
                page.locator('[data-scroll="translation"]').click()
                page.wait_for_function("Math.abs(document.querySelector('#translation').getBoundingClientRect().top - 20) < 2")
                assert page.locator('[data-scroll="translation"]').get_attribute('aria-current') == 'location'
                assert page.locator('#extract-style.secondary').count() == 1
                assert page.locator('[data-file-filter="output"].active').count() == 1
                assert '还没有译文' in page.locator('#file-list').inner_text()
                page.locator('#agent-pill').click()
                page.locator('#project-model-options').select_option('__custom__')
                page.locator('#project-model').fill('test-updated-model')
                page.locator('#save-model').click()
                page.wait_for_function("document.querySelector('#model-status').textContent.includes('已保存：test-updated-model')")
                page.reload()
                page.wait_for_function("document.querySelector('#project-model').value === 'test-updated-model'")
                page.wait_for_function("document.querySelector('#project-model-options').options.length > 2")
                choices = page.locator('#project-model-options option').evaluate_all("els => els.map(el => el.value).filter(v => v && v !== '__custom__')")
                assert choices
                page.locator('#agent-pill').click()
                page.locator('#project-model-options').select_option(choices[0])
                page.locator('#save-model').click()
                page.wait_for_function("document.querySelector('#model-status').textContent.includes('已保存')")
                assert page.locator('#project-model').input_value() == choices[0]
                assert choices[0] in page.locator('#agent-pill').inner_text()
                assert not page.locator('#model-menu').evaluate('el => el.open')
                page.locator('#agent-pill').click()
                page.keyboard.press('Escape')
                assert not page.locator('#model-menu').evaluate('el => el.open')
                page.locator('#style-editor').fill('# Translation Style\n\nUse clear, professional English. Preserve headings and paragraphs.')
                page.wait_for_function("document.querySelector('#style-status').textContent === '已同步'")
                page.locator('[data-history="style"]').click()
                page.locator('#history-dialog').wait_for(state='visible')
                assert page.locator('#history-select option').count() >= 2
                page.locator('#history-select').select_option(index=1)
                page.locator('#restore-history').click()
                page.wait_for_function("document.querySelector('#style-editor').value.includes('Preserve the original meaning')")
                page.locator('#close-history').click()
                doc = Document()
                doc.add_paragraph('The doctor works in a hospital.', style='Heading 1')
                doc.add_paragraph('Keep the second paragraph.')
                buffer = io.BytesIO()
                doc.save(buffer)
                upload = {"name":"产品说明.docx", "mimeType":"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "buffer":buffer.getvalue()}
                page.locator('#source-upload').set_input_files(upload)
                page.wait_for_function("document.querySelector('#source-file').value !== ''")
                page.locator('#corpus-upload').set_input_files({"name":"参考语料.txt", "mimeType":"text/plain", "buffer":"The doctor works in a hospital.\nProfessional documents use clear and precise language.\n医生在医院工作。".encode()})
                page.wait_for_function("document.querySelector('#corpus-summary').textContent.includes('1 份参考语料')")
                page.locator('#corpus-upload').set_input_files([
                    {"name":"追加一.txt","mimeType":"text/plain","buffer":b'The doctor provides medical care to patients.'},
                    {"name":"追加二.txt","mimeType":"text/plain","buffer":b'The hospital provides services to the community.'},
                ])
                page.wait_for_function("document.querySelectorAll('#corpus-list .file-row').length === 3")
                page.once('dialog', lambda dialog: dialog.dismiss())
                page.locator('#corpus-list .file-row').filter(has_text='追加一.txt').locator('button').click()
                assert page.locator('#corpus-list .file-row').count() == 3
                page.once('dialog', lambda dialog: dialog.accept())
                page.locator('#corpus-list .file-row').filter(has_text='追加一.txt').locator('button').click()
                page.wait_for_function("document.querySelectorAll('#corpus-list .file-row').length === 2")
                page.wait_for_function("document.querySelector('#extract-style').classList.contains('primary')")
                page.locator('#extract-style').click()
                page.wait_for_function("document.querySelector('#terms-editor').value.includes('doctor')")
                assert 'doctor' not in page.locator('#style-editor').input_value()
                page.wait_for_function("document.querySelector('#extract-style').classList.contains('secondary')")
                page.locator('#terms-table [data-field="usage"]').fill('Use the preferred term consistently.')
                page.wait_for_function("document.querySelector('#terms-status').textContent === '已同步'")
                page.locator('#tab-mappings').click()
                assert '还没有翻译对照' in page.locator('#mappings-table').inner_text()
                page.locator('[data-add-term="mappings"]').click()
                page.locator('#mappings-table [data-field="original"]').fill('医生')
                page.locator('#mappings-table [data-field="translation"]').fill('doctor')
                page.wait_for_function("document.querySelector('#mappings-status').textContent === '已同步'")
                page.locator('[data-history="mappings"]').click()
                page.locator('#history-dialog').wait_for(state='visible')
                assert page.locator('#history-select option').count() >= 2
                page.locator('#history-select').select_option(index=1)
                page.locator('#restore-history').click()
                page.wait_for_function("JSON.parse(document.querySelector('#mappings-editor').value).rows.length === 0")
                page.locator('#close-history').click()
                page.locator('#tab-terms').click()
                assert page.locator('#terms-table [data-field="usage"]').input_value() == 'Use the preferred term consistently.'

                page.wait_for_function("document.querySelector('#rag-status').textContent.includes('参考索引已同步')")
                assert not page.locator('#build-rag').is_visible()
                assert page.locator('#target-language').input_value() == 'English'
                assert page.locator('#target-language').get_attribute('readonly') is not None
                page.locator('#use-rag').uncheck()
                page.wait_for_function("!document.querySelector('#use-rag').disabled")
                page.reload()
                page.wait_for_function("document.querySelector('#workspace').hidden === false")
                assert not page.locator('#use-rag').is_checked()
                page.locator('#use-rag').check()
                page.wait_for_function("!document.querySelector('#use-rag').disabled")
                page.locator('#recall-query').fill('doctor in hospital')
                page.locator('#test-recall').click()
                page.locator('.recall-result').first.wait_for()
                page.locator('#translate').click()
                page.wait_for_function("document.querySelector('#active-status').dataset.state === 'running' && document.querySelector('#active-title').textContent === '翻译中'")
                assert '1–2' in page.locator('#active-detail').inner_text()
                page.locator('#active-status').screenshot(path='/tmp/transmux-progress.png')
                page.wait_for_function("document.querySelector('#active-title').textContent === '审校中'")
                page.locator('.file-badge').wait_for(timeout=15000)
                with page.expect_download() as download_info:
                    page.locator('.file-row').filter(has=page.locator('.file-badge')).locator('a').click()
                download = download_info.value
                assert Document(download.path()).paragraphs[0].text == '医生在医院工作。'
                page.locator('[data-compare]').first.click()
                page.locator('#comparison-dialog').wait_for(state='visible')
                page.locator('.comparison-pair').first.wait_for()
                workspace_scroll = page.evaluate('comparisonView.workspaceScroll')
                assert page.locator('.comparison-pair').count() == 2
                assert page.locator('.comparison-original .comparison-text').first.inner_text() == 'The doctor works in a hospital.'
                page.locator('[data-revise-paragraph="1"]').click()
                page.locator('[data-revision-form="1"] textarea').fill('Use physician and keep the original meaning.')
                page.locator('[data-revision-form="1"] button[type="submit"]').click()
                page.locator('#comparison-update').wait_for(state='visible', timeout=20000)
                assert page.locator('.comparison-translation .comparison-text').first.inner_text() == '医生在医院工作。'
                assert page.locator('[data-revise-paragraph="2"]').is_disabled()
                position_style = page.add_style_tag(content='.comparison-pair { min-height: 900px; }')
                page.evaluate("""() => {
                    const body = document.querySelector('#comparison-body');
                    const row = document.querySelector('#compare-paragraph-2');
                    body.scrollTop += row.getBoundingClientRect().top - body.getBoundingClientRect().top + 40;
                }""")
                position_before = page.locator('#compare-paragraph-2').evaluate("el => el.getBoundingClientRect().top - document.querySelector('#comparison-body').getBoundingClientRect().top")
                page.locator('#comparison-view-latest').click()
                page.wait_for_function("document.querySelector('.comparison-translation .comparison-text').textContent === 'The physician works at the hospital.'")
                assert page.locator('.comparison-translation .comparison-text').nth(1).inner_text() == '保留第二段。'
                position_after = page.locator('#compare-paragraph-2').evaluate("el => el.getBoundingClientRect().top - document.querySelector('#comparison-body').getBoundingClientRect().top")
                assert abs(position_after - position_before) <= 2
                position_style.evaluate('el => el.remove()')
                assert page.locator('#comparison-version option').count() == 2
                with page.expect_download() as revision_download:
                    page.locator('#comparison-download').click()
                revised_doc = Document(revision_download.value.path())
                assert revised_doc.paragraphs[0].text == 'The physician works at the hospital.'
                assert revised_doc.paragraphs[1].text == '保留第二段。'
                page.locator('#comparison-dialog').screenshot(path='/tmp/transmux-comparison.png')
                page.set_viewport_size({"width":390,"height":844})
                assert page.locator('.comparison-side-label').first.is_visible()
                assert page.locator('#comparison-dialog').evaluate('el => el.scrollWidth <= el.clientWidth')
                page.locator('#comparison-dialog').screenshot(path='/tmp/transmux-comparison-mobile.png')
                page.locator('#close-comparison').click()
                assert abs(page.evaluate('window.scrollY') - workspace_scroll) <= 1
                page.set_viewport_size({"width":1440,"height":1050})
                page.locator('[data-compare]').first.click()
                page.wait_for_function("document.querySelectorAll('#comparison-version option').length === 2")
                page.keyboard.press('Escape')
                # Open a translated result, then exercise the independent DOCX entry.
                page.locator('[data-layout]').first.click()
                page.locator('#layout-dialog').wait_for(state='visible')
                assert page.locator('#layout-template').input_value() == 'original'
                page.locator('#close-layout').click()
                source_count = page.locator('#source-file option').count()
                page.locator('.layout-entry summary').click()
                page.locator('#manuscript-upload').set_input_files({**upload, 'name':'独立文稿.docx'})
                page.locator('#layout-dialog').wait_for(state='visible')
                assert page.locator('#layout-source').inner_text() == '独立文稿.docx'
                assert page.locator('#source-file option').count() == source_count
                with page.expect_download() as original_download:
                    page.locator('#layout-original').click()
                assert open(original_download.value.path(), 'rb').read() == buffer.getvalue()
                page.locator('#layout-generate').click()
                page.locator('#layout-pdf').wait_for(state='visible')
                page.wait_for_function("document.querySelector('#layout-checks').textContent.includes('字节完全一致')")
                first_version = page.locator('#layout-version').input_value()
                first_preview = page.locator('#layout-pdf').get_attribute('src')
                assert page.request.get('http://127.0.0.1:18765' + first_preview).headers['content-type'] == 'application/pdf'
                with page.expect_download() as export_download:
                    page.locator('#layout-docx').click()
                assert open(export_download.value.path(), 'rb').read() == buffer.getvalue()
                page.locator('#layout-generate').click()
                page.wait_for_function("document.querySelectorAll('#layout-version option').length === 2")
                assert page.locator('#layout-version').input_value() != first_version
                page.locator('#layout-version').select_option(first_version)
                assert page.locator('#layout-pdf').get_attribute('src') == first_preview
                page.locator('#layout-template').select_option('ieee-access')
                page.wait_for_function("document.querySelectorAll('[data-layout-role]').length === 2")
                page.locator('#layout-tab-structure').click()
                page.locator('[data-layout-role]').first.select_option('title')
                page.locator('#layout-generate').click()
                page.wait_for_function("document.querySelectorAll('#layout-version option').length === 3")
                page.wait_for_function("document.querySelector('#layout-checks').textContent.includes('投稿排版草稿')")
                page.locator('#layout-tab-report').click()
                assert '原样保留' in page.locator('#layout-checks').inner_text()
                assert page.locator('#layout-status').get_attribute('data-state') == 'needs_attention'
                with page.expect_download() as preset_download:
                    page.locator('#layout-docx').click()
                preset_doc = Document(preset_download.value.path())
                assert len(preset_doc.sections) == 2
                assert preset_doc.paragraphs[0].style.font.size.pt == 22
                assert preset_doc.paragraphs[0].text == doc.paragraphs[0].text
                page.locator('#layout-template').select_option('jcst')
                page.locator('#layout-tab-structure').click()
                assert page.locator('[data-layout-role]').first.input_value() == 'title'
                structure_box = page.locator('#layout-structure-list').bounding_box()
                assert structure_box['width'] > 1300 and structure_box['height'] > 600, structure_box
                page.locator('#layout-dialog').screenshot(path='/tmp/transmux-layout-structure.png')
                # Switching tabs keeps structure overrides and does not reload the PDF.
                page.locator('#layout-tab-preview').click()
                assert page.locator('#layout-pdf').get_attribute('src') != first_preview
                page.locator('#layout-tab-preview').focus()
                page.keyboard.press('ArrowRight')
                assert page.locator('#layout-structure').is_visible()
                page.locator('#layout-generate').click()
                page.wait_for_function("document.querySelectorAll('#layout-version option').length === 4")
                assert page.locator('#layout-preview').is_visible()
                assert 'JCST' in page.locator('#layout-version option:checked').inner_text()
                with page.expect_download() as jcst_download:
                    page.locator('#layout-docx').click()
                jcst_doc = Document(jcst_download.value.path())
                assert jcst_doc.paragraphs[0].style.font.size.pt == 16
                assert jcst_doc.sections[-1].left_margin.twips == 839
                pdf_box = page.locator('#layout-pdf').bounding_box()
                assert pdf_box['width'] >= 1430 and pdf_box['height'] > 650, pdf_box
                page.locator('#layout-tab-report').click()
                assert 'JCST' in page.locator('#layout-checks').inner_text()
                assert 'sciopen.com' in page.locator('#layout-checks > a').get_attribute('href')
                page.locator('#layout-tab-preview').click()
                page.locator('#layout-dialog').screenshot(path='/tmp/transmux-layout.png')
                page.set_viewport_size({'width':390, 'height':844})
                assert page.locator('#layout-dialog').evaluate('el => el.scrollWidth <= el.clientWidth')
                assert page.locator('#layout-pdf').bounding_box()['width'] >= 380
                page.locator('#layout-dialog').screenshot(path='/tmp/transmux-layout-mobile.png')
                page.locator('#layout-tab-structure').click()
                assert page.locator('[data-layout-role]').first.is_visible()
                assert page.locator('#layout-structure-list').bounding_box()['height'] > 250
                assert page.locator('#layout-dialog').evaluate('el => el.scrollWidth <= el.clientWidth')
                page.locator('#layout-dialog').screenshot(path='/tmp/transmux-layout-structure-mobile.png')
                page.locator('#layout-template').select_option('original')
                assert page.locator('#layout-preview').is_visible()
                assert page.locator('#layout-tab-structure').is_hidden()
                page.keyboard.press('Escape')
                assert not page.locator('body').evaluate('el => el.classList.contains("layout-open")')
                page.set_viewport_size({'width':1440, 'height':1050})
                page.locator('#chat-input').fill('检查文件并告诉我当前状态')
                page.locator('#chat-input').press('Enter')
                page.wait_for_function("document.querySelector('#messages').textContent.includes('检查文件并告诉我当前状态')")
                page.evaluate('window.scrollTo(0, 0)')
                page.screenshot(path='/tmp/transmux-workspace.png', full_page=True)
                page.set_viewport_size({"width":390,"height":844})
                page.locator('#agent-pill').click()
                assert page.locator('.model-popover').is_visible()
                assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
                page.keyboard.press('Escape')
                page.screenshot(path='/tmp/transmux-mobile.png', full_page=True)
                page.wait_for_function("!document.querySelector('#delete-workspace').disabled")
                page.locator('.workspace-settings summary').click()
                page.locator('#delete-workspace').click()
                page.locator('#delete-workspace-confirm').fill('错误名称')
                assert page.locator('#confirm-delete-workspace').is_disabled()
                page.locator('#cancel-delete-workspace').click()
                assert page.locator('#workspace').is_visible()
                page.locator('#delete-workspace').click()
                page.locator('#delete-workspace-confirm').fill('产品文档 · 中英翻译')
                page.locator('#confirm-delete-workspace').click()
                page.locator('#onboarding').wait_for(state='visible')
                assert page.locator('#project-list .project-button').count() == 0
                page.reload()
                page.locator('#onboarding').wait_for(state='visible')
                assert not errors, errors
                browser.close()
                print('Browser smoke passed: onboarding, autosave, upload, RAG, recall, translation, download, chat, mobile; no JS errors.')
        finally:
            server.should_exit = True
            thread.join(timeout=10)


if __name__ == '__main__':
    main()
