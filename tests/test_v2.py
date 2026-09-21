import asyncio
import io
import json
import time

from docx import Document
from fastapi.testclient import TestClient

from transmux.citations import apply_citation_edits
from transmux.documents import extract
from transmux.v2 import WorkspaceStore, WorkspaceWorker, create_v2_app
from test_workflows import FakeRunner


class Runner(FakeRunner):
    async def run_isolated(self, pid, jid, prompt, schema=None):
        self.calls.append((pid, jid, prompt, schema))
        properties = schema['properties'] if schema else {}
        if 'findings' in properties:
            if getattr(self, 'store', None):
                self.store.event(pid, jid, 'agent', json.dumps({'type': 'item.completed', 'item': {'type': 'web_search'}}))
            return {'findings': [{'paragraph': 1, 'quote': 'Paris is in Germany.', 'status': 'incorrect',
                'explanation': 'Paris is the capital of France.', 'replacement': 'Paris is in France.',
                'sources': [{'url': 'https://www.paris.fr/', 'title': 'Paris', 'evidence': 'Official municipal website in France.'}]}]}
        if 'answer' in properties:
            return {'answer': 'Hello from your workspace.', 'edits': []}
        if 'edits' in properties:
            return {'edits': [], 'issues': []}
        if 'translations' in properties:
            # Fixture has one paragraph; real alignment/review/export still execute.
            return {'translations': [{'source_ids': ['p000001'], 'paragraphs': ['Paris is in Germany.'], 'reason': ''}], 'mappings': [], 'people': []}
        if 'passed' in properties:
            return {'passed': True, 'issues': [], 'mapping_issues': [], 'people_issues': []}
        return await super().run(pid, jid, prompt, schema)


def upload(client, pid, text='Paris is in Germany.', name='paper.docx'):
    doc = Document()
    doc.add_paragraph(text)
    stream = io.BytesIO()
    doc.save(stream)
    response = client.post(f'/api/projects/{pid}/attachments', files={'file': (name, stream.getvalue())})
    assert response.status_code == 201, response.text
    return response.json()


def finished(client, pid, jid):
    for _ in range(200):
        jobs = client.get(f'/api/projects/{pid}/jobs').json()
        row = next(j for j in jobs if j['id'] == jid)
        if row['state'] not in ('queued', 'running'):
            return row
        time.sleep(.05)
    raise AssertionError('job did not finish')


def test_fresh_workspace_scope_fact_followup_and_versions(tmp_path, monkeypatch):
    monkeypatch.setattr('transmux.app.availability', lambda: [{'id': 'codex', 'available': True}, {'id': 'codebuddy', 'available': True}])
    runner = Runner()
    def factory(store):
        runner.store = store
        return WorkspaceWorker(store, runner)
    app = create_v2_app(tmp_path, factory)
    with TestClient(app) as client:
        assert client.get('/api/projects').json() == []
        assert client.get('/api/health').json()['embedding'] is None
        pid = client.post('/api/projects', json={'name': 'New', 'agent': 'codex'}).json()['id']
        other = client.post('/api/projects', json={'name': 'Other', 'agent': 'codebuddy'}).json()['id']
        source = upload(client, pid)
        upload(client, pid, 'This other upload must not be processed.', 'unselected.docx')
        assert client.post(f'/api/projects/{pid}/messages', json={'kind': 'translate'}).status_code == 400
        assert client.post(f'/api/projects/{other}/messages', json={'kind': 'factcheck', 'file_ids': [source['id']]}).status_code == 400
        assert client.post(f'/api/projects/{pid}/jobs', json={'kind': 'rag'}).status_code == 404
        assert not (tmp_path / 'projects' / pid / 'rag').exists()
        task = client.post(f'/api/projects/{pid}/messages', json={'kind': 'translate', 'file_ids': [source['id']], 'message': 'Use concise English.'}).json()
        result = finished(client, pid, task['id'])
        assert result['state'] == 'succeeded', result
        translated = json.loads(result['result'])['documents'][0]['file_id']
        assert len([f for f in client.get(f'/api/projects/{pid}/files').json() if f['kind'] == 'output']) == 1
        assert any('Use concise English.' in call[2] and 'transmux-translate' in call[2] for call in runner.calls)
        check = client.post(f'/api/projects/{pid}/messages', json={'kind': 'factcheck'}).json()
        assert json.loads(check['payload'])['file_ids'] == [translated]
        checked = finished(client, pid, check['id'])
        assert checked['state'] == 'succeeded', checked
        original = client.get(f'/api/projects/{pid}/files/{translated}/download').content
        result = json.loads(checked['result'])
        preview = client.get(f"/api/projects/{pid}/files/{result['file_id']}/preview").json()
        assert 'https://www.paris.fr/' in preview['content']
        assert client.post(f'/api/projects/{pid}/messages', json={'message': '是', 'report_job': 'stale'}).status_code == 409
        correction = client.post(f'/api/projects/{pid}/messages', json={'message': '是', 'report_job': check['id']}).json()
        assert correction['kind'] == 'factfix'
        corrected = finished(client, pid, correction['id'])
        assert corrected['state'] == 'succeeded', corrected
        fid = json.loads(corrected['result'])['file_id']
        content = client.get(f'/api/projects/{pid}/files/{fid}/preview').json()
        assert content['paragraphs'] == ['Paris is in France.']
        assert client.get(f'/api/projects/{pid}/files/{translated}/download').content == original
        again = client.post(f'/api/projects/{pid}/messages', json={'message': '是'}).json()
        assert again['kind'] == 'chat'
        assert client.get(f'/api/projects/{other}/files/{fid}/preview').status_code == 400


def test_style_only_enrolls_selected_reference_and_caches_without_rag(tmp_path):
    store = WorkspaceStore(tmp_path)
    pid = store.create_project('References', 'codex')['id']
    runner = Runner()
    worker = WorkspaceWorker(store, runner)
    ids = []
    for name, text in [('selected', 'The doctor works in a hospital. We describe each method clearly and avoid unsupported claims.'), ('other', 'This unselected document must not affect the style guide.')]:
        path = store.workspace(pid) / 'sources' / (name + '.txt')
        path.write_text(text)
        path.with_suffix('.txt.json').write_text(json.dumps([text]))
        ids.append(store.add_file(pid, path.name, 'source', path))
    async def run():
        for _ in range(2):
            job = store.enqueue(pid, 'style', {'file_ids': [ids[0]], 'use_rag': False})
            await worker.execute(job)
            result = store.rows('SELECT * FROM jobs WHERE id=?', (job['id'],))[0]
            assert result['state'] == 'succeeded', result
    asyncio.run(run())
    assert store.file(pid, ids[1])['kind'] == 'source'
    assert not any('unselected document' in call[2] for call in runner.calls)
    assert not (store.workspace(pid) / 'rag').exists()
    assert len(list((store.workspace(pid) / 'style-cache').glob('*.json'))) > 0
    references = worker.corpus(pid)
    assert len(references) == 1
    store.delete_corpus(pid, references[0]['id'])
    assert worker.corpus(pid) == []
    assert store.download_path(pid, store.file(pid, ids[0])).is_file()
    store.db.close()


def test_citation_edits_do_not_rewrite_body_or_invent_metadata(tmp_path):
    path = tmp_path / 'paper.docx'
    doc = Document()
    doc.add_paragraph('A finding (Smith, 2020).')
    doc.add_paragraph('Smith. A title. Journal of Tests. 2020.')
    doc.save(path)
    applied, issues = apply_citation_edits(path, [
        {'paragraph': 1, 'original': 'A finding', 'replacement': 'A different conclusion', 'italic': []},
        {'paragraph': 1, 'original': '(Smith, 2020)', 'replacement': '[1]', 'italic': []},
        {'paragraph': 2, 'original': 'Smith. A title. Journal of Tests. 2020.', 'replacement': 'Smith. A title. Journal of Tests. 2020. DOI invented.', 'italic': []},
        {'paragraph': 2, 'original': 'Smith. A title. Journal of Tests. 2020.', 'replacement': 'Smith, A title, Journal of Tests, 2020.', 'italic': ['Journal of Tests']},
    ], {2})
    assert applied == 2 and len(issues) == 2
    assert extract(path)[0] == 'A finding [1].'
    assert any(r.italic and r.text == 'Journal of Tests' for r in Document(path).paragraphs[1].runs)


def test_fact_report_without_observed_external_tools_cannot_correct(tmp_path):
    store = WorkspaceStore(tmp_path)
    pid = store.create_project('Unverified', 'codebuddy')['id']
    path = store.workspace(pid) / 'sources' / 'claim.txt'
    path.write_text('Paris is in Germany.')
    fid = store.add_file(pid, path.name, 'source', path)
    worker = WorkspaceWorker(store, Runner())  # No simulated external-tool event.
    task = store.enqueue(pid, 'factcheck', {'file_ids': [fid], 'use_rag': False})
    asyncio.run(worker.execute(task))
    report = json.loads((store.workspace(pid) / 'runs' / task['id'] / 'factcheck.json').read_text())
    assert report['findings'][0]['status'] == 'insufficient'
    assert report['findings'][0]['replacement'] == ''
    store.db.close()


def test_layout_artifacts_are_previewable_and_original_format_is_unchanged(tmp_path, monkeypatch):
    from test_layout import fake_pdf
    monkeypatch.setattr('transmux.layout.render_pdf', fake_pdf)
    store = WorkspaceStore(tmp_path)
    pid = store.create_project('Layout', 'codex')['id']
    path = store.workspace(pid) / 'sources' / 'paper.docx'
    doc = Document()
    doc.add_heading('A Paper', 0)
    doc.add_heading('Abstract', 1)
    doc.add_paragraph('We describe a reproducible method.')
    doc.add_heading('1 Introduction', 1)
    doc.add_paragraph('This body must remain unchanged.')
    doc.save(path)
    fid = store.add_file(pid, path.name, 'source', path)
    worker = WorkspaceWorker(store, Runner())
    async def run():
        for template in ('original', 'jcst', 'ieee-access'):
            task = store.enqueue(pid, 'layout', {'file_ids': [fid], 'template': template, 'use_rag': False})
            await worker.execute(task)
            row = store.rows('SELECT state,result FROM jobs WHERE id=?', (task['id'],))[0]
            assert row['state'] in ('succeeded', 'needs_attention'), row
            result = json.loads(row['result'])
            files = [store.file(pid, id) for id in result['file_ids']]
            docx = next(f for f in files if f['name'].endswith('.docx'))
            output = store.download_path(pid, docx)
            assert extract(output) == extract(path)
            if template == 'original':
                assert output.read_bytes() == path.read_bytes()
            assert any(f['name'].endswith('.pdf') for f in files)
    asyncio.run(run())
    store.db.close()


def test_chat_formatting_creates_a_new_version_without_changing_text(tmp_path):
    class FormatRunner(Runner):
        async def run_isolated(self, pid, jid, prompt, schema=None):
            data = json.loads(prompt.split('\n')[-1])
            return {'answer': 'Formatted the selected document.', 'edits': [], 'formatting': [{
                'file_id': data['files'][0]['id'], 'font': 'Times New Roman', 'font_size': 12,
                'columns': 2, 'line_spacing': 1.5}]}
    store = WorkspaceStore(tmp_path)
    pid = store.create_project('Editing', 'codex')['id']
    path = store.workspace(pid) / 'sources' / 'paper.docx'
    doc = Document()
    p = doc.add_paragraph('Keep ')
    p.add_run('this text').italic = True
    doc.save(path)
    original = path.read_bytes()
    fid = store.add_file(pid, path.name, 'source', path)
    worker = WorkspaceWorker(store, FormatRunner())
    task = store.enqueue(pid, 'chat', {'file_ids': [fid], 'message': 'Use 12pt Times New Roman, double columns and 1.5 line spacing.'})
    asyncio.run(worker.execute(task))
    row = store.rows('SELECT state,result FROM jobs WHERE id=?', (task['id'],))[0]
    assert row['state'] == 'succeeded', row
    output = store.file(pid, json.loads(row['result'])['file_ids'][0])
    result = Document(store.download_path(pid, output))
    assert result.paragraphs[0].text == 'Keep this text'
    assert result.paragraphs[0].runs[-1].italic
    assert result.paragraphs[0].runs[0].font.size.pt == 12
    assert path.read_bytes() == original
    store.db.close()


def test_original_doc_preview_uses_converted_copy_and_keeps_download(tmp_path, monkeypatch):
    app = create_v2_app(tmp_path, lambda store: WorkspaceWorker(store, Runner()))
    with TestClient(app) as client:
        # Seed inside the application thread so SQLite retains its ownership invariant.
        @app.get('/fixture-original-doc')
        async def fixture():
            store = app.state.store
            pid = store.create_project('DOC preview', 'codex')['id']
            path = store.workspace(pid) / 'sources' / 'original.doc'
            path.write_bytes(b'original-binary-doc')
            converted = Document()
            converted.add_paragraph('Converted reading content.')
            converted.save(path.with_suffix('.docx'))
            fid = store.add_file(pid, path.name, 'original', path)
            return {'pid': pid, 'fid': fid}
        # The static mount consumes unmatched routes; move this test-only fixture before it.
        app.router.routes.insert(0, app.router.routes.pop())
        ids = client.get('/fixture-original-doc').json()
        url = f"/api/projects/{ids['pid']}/files/{ids['fid']}"
        assert client.get(url + '/preview').json()['paragraphs'] == ['Converted reading content.']
        assert client.get(url + '/download').content == b'original-binary-doc'


def test_event_tail_is_bounded_ordered_and_project_scoped(tmp_path, monkeypatch):
    monkeypatch.setattr('transmux.app.availability', lambda: [{'id': 'codex', 'available': True}])
    app = create_v2_app(tmp_path, lambda store: WorkspaceWorker(store, Runner()))
    with TestClient(app) as client:
        pid = client.post('/api/projects', json={'name': 'Events', 'agent': 'codex'}).json()['id']
        other = client.post('/api/projects', json={'name': 'Other', 'agent': 'codex'}).json()['id']
        for i in range(350):
            client.portal.call(app.state.store.event, pid, None, 'progress', str(i))
        client.portal.call(app.state.store.event, other, None, 'progress', 'private')
        rows = client.get(f'/api/projects/{pid}/events?tail=true').json()
        assert len(rows) == 300
        assert [r['text'] for r in rows] == [str(i) for i in range(50, 350)]
        assert client.get(f'/api/projects/{pid}/events').json()[0]['text'] == '0'
        client.portal.call(app.state.store.event, pid, None, 'progress', 'new')
        assert [r['text'] for r in client.get(f"/api/projects/{pid}/events?after={rows[-1]['id']}").json()] == ['new']


def test_translation_does_not_preload_terminology_or_names(tmp_path, monkeypatch):
    from transmux import terminology
    monkeypatch.setattr('transmux.app.availability', lambda: [{'id': 'codex', 'available': True}])
    runner = Runner()
    app = create_v2_app(tmp_path, lambda store: WorkspaceWorker(store, runner))
    with TestClient(app) as client:
        pid = client.post('/api/projects', json={'name': 'Style only', 'agent': 'codex'}).json()['id']
        source = upload(client, pid)
        glossary = upload(client, pid, 'ATTACHED_GLOSSARY_SENTINEL', 'glossary.docx')
        saved = {}
        for kind in terminology.KINDS:
            rows = [{**{field: 'REGISTRY_SENTINEL ' + str(i) for field in terminology.FIELDS[kind]},
                     'origin': 'manual', 'context' if kind != 'terms' else 'usage': 'Long registry evidence. ' * 200}
                    for i in range(50)]
            content = terminology.encode({'rows': rows, 'deleted': [], 'legacy': ''})
            def write(kind=kind, content=content):
                store = app.state.store
                store.write_config(pid, kind, content, store.snapshot_config(pid, kind)['revision'])
                return store.snapshot_config(pid, kind)['content']
            saved[kind] = client.portal.call(write)
        style = client.get(f'/api/projects/{pid}/config/style').json()['content']
        task = client.post(f'/api/projects/{pid}/messages', json={'kind': 'translate', 'file_ids': [source['id']],
                          'glossary_ids': [glossary['id']], 'message': 'Keep the original meaning.'}).json()
        result = finished(client, pid, task['id'])
        assert result['state'] == 'succeeded', result['result']
        assert len(runner.calls) == 2  # translation and review both exclude registries
        for _, _, prompt, _ in runner.calls:
            assert style in prompt and 'Keep the original meaning.' in prompt
            assert 'REGISTRY_SENTINEL' not in prompt
            assert 'ATTACHED_GLOSSARY_SENTINEL' not in prompt
            assert 'Check names against approved spellings' not in prompt
        for kind, content in saved.items():
            assert client.get(f'/api/projects/{pid}/config/{kind}').json()['content'] == content
