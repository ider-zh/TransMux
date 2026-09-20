import asyncio
import hashlib
import io
import json
import sqlite3
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
from fastapi.testclient import TestClient
from pypdf import PdfWriter
import pytest

from transmux.app import create_app
from transmux.jobs import Worker
from transmux.layout import export_path, inspect_docx, render_pdf
from transmux.rag import Rag
from transmux.store import Store
from test_comparison import wait_job
from test_workflows import FakeRunner, TinyEmbeddings


def manuscript():
    doc = Document()
    doc.add_heading('Document with structure', 1)
    paragraph = doc.add_paragraph('Original citation [1]. ')
    run = paragraph.add_run('Italic text')
    run.italic = True
    run.font.size = Pt(13)
    paragraph._p.append(OxmlElement('m:oMath'))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Merged header'
    table.cell(1, 0).text = 'Value'
    doc.sections[0].header.paragraphs[0].text = 'Running header'
    doc.sections[0].footer.paragraphs[0].text = 'Footer'
    result = io.BytesIO()
    doc.save(result)
    return result.getvalue()


async def fake_pdf(docx, directory):
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    with (directory / 'document.pdf').open('wb') as stream:
        writer.write(stream)
    return {'engine':'test renderer', 'pages':1}


def worker(store):
    return Worker(store, FakeRunner(), Rag(TinyEmbeddings()))


def test_export_api_history_original_bytes_isolation_restart_and_deletion(tmp_path, monkeypatch):
    monkeypatch.setattr('transmux.layout.render_pdf', fake_pdf)
    monkeypatch.setattr('transmux.app.availability', lambda: [{'id':'codex','available':True}])
    app = create_app(tmp_path, worker)
    original = manuscript()
    with TestClient(app) as client:
        pid = client.post('/api/projects', json={'name':'layout', 'agent':'codex'}).json()['id']
        other = client.post('/api/projects', json={'name':'other', 'agent':'codex'}).json()['id']
        base = f'/api/projects/{pid}'
        uploaded = client.post(base + '/files?kind=manuscript', files={'file':('paper.docx', original)})
        assert uploaded.status_code == 201
        fid = uploaded.json()['id']
        assert client.get(base + '/files').json()[0]['can_layout']
        assert client.post(base + '/files?kind=manuscript', files={'file':('paper.txt', b'Text')}).status_code == 400
        assert client.post(base + '/jobs', json={'kind':'layout','file_id':fid,'template':'ieee'}).status_code == 422
        assert client.post(f'/api/projects/{other}/jobs', json={'kind':'layout','file_id':fid}).status_code == 400
        for _ in range(2):
            task = client.post(base + '/jobs', json={'kind':'layout','file_id':fid}).json()
            result = wait_job(client, base, task['id'])
            assert result['state'] == 'succeeded', result
        versions = client.get(base + f'/files/{fid}/exports').json()
        assert len(versions) == 2
        assert versions[0]['id'] != versions[1]['id']
        for version in versions:
            eid = version['id']
            manifest = version['manifest']
            assert manifest['source_sha256'] == hashlib.sha256(original).hexdigest()
            assert manifest['checks'][0]['passed']
            assert client.get(base + f'/exports/{eid}/docx').content == original
            pdf = client.get(base + f'/exports/{eid}/pdf?preview=true')
            assert pdf.content.startswith(b'%PDF')
            assert pdf.headers['content-disposition'].startswith('inline')
            assert pdf.headers['x-frame-options'] == 'SAMEORIGIN'
            assert client.get(base + f'/exports/{eid}/pdf').headers['content-disposition'].startswith('attachment')
            assert client.get(base + f'/exports/{eid}/json').json() == manifest
            assert client.get(f'/api/projects/{other}/exports/{eid}/docx').status_code == 400
        assert client.get(base + f'/files/{fid}/download').content == original
        assert client.get(base + f'/files/{fid}/exports').headers['x-frame-options'] == 'DENY'
        assert app.state.worker.runner.calls == []
    with TestClient(app) as client:
        assert len(client.get(base + f'/files/{fid}/exports').json()) == 2
        assert client.request('DELETE', base, json={'name':'layout'}).status_code == 200
        assert not (tmp_path / 'exports' / pid).exists()
        with sqlite3.connect(tmp_path / 'transmux.sqlite3') as db:
            assert db.execute('SELECT * FROM layout_exports').fetchall() == []


async def test_pdf_failure_keeps_docx_and_reports_attention_without_agent(tmp_path, monkeypatch):
    store = Store(tmp_path)
    pid = store.create_project('layout', 'codex')['id']
    source = store.workspace(pid) / 'sources' / 'paper.docx'
    source.write_bytes(manuscript())
    fid = store.add_file(pid, 'paper.docx', 'manuscript', source)

    async def fail(*args):
        raise ValueError('PDF 转换失败')

    monkeypatch.setattr('transmux.layout.render_pdf', fail)
    runner = worker(store)
    task = store.enqueue(pid, 'layout', {'file_id':fid})
    await runner.execute(task)
    result = store.rows('SELECT * FROM jobs WHERE id=?', (task['id'],))[0]
    assert result['state'] == 'needs_attention'
    eid = json.loads(result['result'])['export_id']
    output, record = export_path(store, pid, eid, 'docx')
    assert output.read_bytes() == source.read_bytes()
    assert record['manifest']['issues'] == ['PDF 转换失败']
    assert not record['manifest']['pdf']
    with pytest.raises(ValueError, match='PDF'):
        export_path(store, pid, eid, 'pdf')
    assert runner.runner.calls == []
    store.db.close()


async def test_cancel_export_never_publishes_partial_version(tmp_path, monkeypatch):
    store = Store(tmp_path)
    pid = store.create_project('layout', 'codex')['id']
    source = store.workspace(pid) / 'sources' / 'paper.docx'
    original = manuscript()
    source.write_bytes(original)
    fid = store.add_file(pid, 'paper.docx', 'manuscript', source)
    ready = asyncio.Event()

    async def wait(*args):
        ready.set()
        await asyncio.Event().wait()

    monkeypatch.setattr('transmux.layout.render_pdf', wait)
    task = asyncio.create_task(worker(store).execute(store.enqueue(pid, 'layout', {'file_id':fid})))
    await ready.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert store.rows('SELECT * FROM layout_exports') == []
    assert list((tmp_path / 'exports' / pid).iterdir()) == []
    assert source.read_bytes() == original
    store.db.close()


def test_linked_resources_and_fields_are_not_rendered(tmp_path):
    source = tmp_path / 'paper.docx'
    source.write_bytes(manuscript())
    assert inspect_docx(source) == []
    with ZipFile(source, 'a') as archive:
        archive.writestr('word/_rels/extra.xml.rels', '<Relationships><Relationship TargetMode="External" Type="image" Target="https://example.invalid/image"/></Relationships>')
    assert '外部链接资源' in inspect_docx(source)[0]
    field_doc = Document()
    field = OxmlElement('w:instrText')
    field.text = ' INCLUDETEXT "https://example.invalid/data" '
    field_doc.add_paragraph()._p.append(field)
    field_doc.save(source)
    assert '外部数据字段' in inspect_docx(source)[0]


async def test_renderer_missing_invalid_output_and_process_cleanup(tmp_path, monkeypatch):
    monkeypatch.setattr('transmux.layout.converter_command', lambda: None)
    source = tmp_path / 'document.docx'
    source.write_bytes(manuscript())
    with pytest.raises(ValueError, match='LibreOffice'):
        await render_pdf(source, tmp_path)

    # Exercise the actual process adapter without relying on an installed renderer.
    executable = tmp_path / 'fake-converter'
    executable.write_text('#!/usr/bin/env python3\nfrom pathlib import Path\nimport sys\nPath(sys.argv[sys.argv.index("--outdir")+1], "document.pdf").write_bytes(b"bad PDF")\n')
    executable.chmod(0o700)
    monkeypatch.setattr('transmux.layout.converter_command', lambda: str(executable))
    with pytest.raises(ValueError, match='有效 PDF'):
        await render_pdf(source, tmp_path)
    assert not (tmp_path / 'profile').exists()

    marker = tmp_path / 'pid'
    executable.write_text(f'#!/usr/bin/env python3\nimport os,time\nfrom pathlib import Path\nPath({str(marker)!r}).write_text(str(os.getpid()))\ntime.sleep(60)\n')
    task = asyncio.create_task(render_pdf(source, tmp_path))
    for _ in range(100):
        if marker.exists():
            break
        await asyncio.sleep(.02)
    assert marker.exists()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert not Path('/proc/' + marker.read_text()).exists()
    assert not (tmp_path / 'profile').exists()
