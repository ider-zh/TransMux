import asyncio
import json
import sqlite3
from pathlib import Path

import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient

from transmux.agents import command, parse_json
from transmux.app import create_app
from transmux.documents import extract, export_docx
from transmux.jobs import Worker
from transmux.rag import Rag
from transmux.store import Store


class TinyEmbeddings:
    identity = "test-semantic-model"

    def encode(self, texts):
        vectors = []
        for text in texts:
            health = any(word in text.lower() for word in ("doctor", "医生", "医院", "hospital"))
            vectors.append([1, 0] if health else [0, 1])
        return np.array(vectors, dtype=np.float32)


class FakeRunner:
    def __init__(self, passed=True):
        self.passed = passed
        self.calls = []

    async def run(self, pid, jid, prompt, schema=None):
        self.calls.append((pid, jid, prompt, schema))
        if schema is None:
            return "已完成"
        if "observations" in schema["properties"]:
            paragraphs = json.loads(prompt.split('\n')[-1])['paragraphs']
            sample = next((r for r in paragraphs if r['style_sample']), None)
            return {'observations': [dict(rule='Use clear, professional English.', paragraph=sample['paragraph'], quote=sample['text'])] if sample else [],
                    'terms': [dict(term='doctor', meaning='A medical professional', usage='Use in a medical context', paragraph=1,
                                   need='preferred_variant', scope='Clinical terminology', reason='Use doctor consistently instead of physician')], 'people': []}
        if "rules" in schema["properties"]:
            return {'rules': [dict(text='Use clear, professional English.', evidence_ids=[1])]}
        if "style" in schema["properties"]:
            return {"style": "# Updated Style\nUse clear, professional English.", "terms": [{"term":"doctor", "meaning":"A medical professional", "usage":"Use in a medical context", "paragraph":1, "need":"preferred_variant", "scope":"Clinical terminology", "reason":"Use doctor consistently instead of alternating with physician in this project"}]}
        if "translations" in schema["properties"]:
            return {"translations": ["医生在医院工作。", "保留第二段。"]}
        return {"passed": self.passed, "issues": [] if self.passed else ["术语不一致"]}


@pytest.fixture
def store(tmp_path):
    instance = Store(tmp_path)
    yield instance
    instance.db.close()


def add_source(store, pid, kind="source"):
    work = store.workspace(pid)
    source = work / ("sources" if kind == "source" else "corpus") / "sample.docx"
    doc = Document()
    doc.add_paragraph("The doctor works in a hospital.", style="Heading 1")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Keep the second paragraph."
    doc.save(source)
    Path(str(source) + ".json").write_text(json.dumps(extract(source)))
    return store.add_file(pid, "sample.docx", kind, source)


def job(store, pid, fid, rounds=3):
    return store.enqueue(pid, "translate", {"file_id": fid, "target_language": "简体中文", "max_review_rounds": rounds})


def test_binding_and_project_isolation(store):
    a = store.create_project("A", "codex")
    b = store.create_project("B", "codebuddy")
    with pytest.raises(sqlite3.IntegrityError, match="immutable"):
        store.execute("UPDATE projects SET agent='codebuddy' WHERE id=?", (a["id"],))
    fid = add_source(store, a["id"])
    with pytest.raises(ValueError):
        store.file(b["id"], fid)
    with pytest.raises(ValueError):
        store.safe_path(a["id"], "../../../transmux.sqlite3")


async def test_translate_review_and_downloadable_docx(store):
    pid = store.create_project("A", "codex")["id"]
    fid = add_source(store, pid)
    runner = FakeRunner()
    worker = Worker(store, runner, Rag(TinyEmbeddings()))
    queued = job(store, pid, fid)
    await worker.execute(queued)
    result = store.rows("SELECT * FROM jobs WHERE id=?", (queued["id"],))[0]
    assert result["state"] == "succeeded"
    file = store.file(pid, json.loads(result["result"])["file_id"])
    output = store.safe_path(pid, file["path"])
    doc = Document(output)
    assert doc.paragraphs[0].text == "医生在医院工作。"
    assert doc.paragraphs[0].style.name == "Heading 1"
    assert doc.tables[0].cell(0, 0).text == "保留第二段。"
    assert len(runner.calls) == 2


async def test_failed_review_never_publishes_and_retries(store):
    pid = store.create_project("A", "codex")["id"]
    runner = FakeRunner(passed=False)
    worker = Worker(store, runner, Rag(TinyEmbeddings()))
    queued = job(store, pid, add_source(store, pid), rounds=2)
    await worker.execute(queued)
    assert store.rows("SELECT state FROM jobs")[0]["state"] == "needs_attention"
    assert not store.rows("SELECT * FROM files WHERE kind='output'")
    assert len(runner.calls) == 4
    assert "术语不一致" in runner.calls[2][2]
    assert len(list((store.workspace(pid) / "runs" / queued["id"]).glob('batch-1-review-*.json'))) == 2


async def test_rag_required_and_references_injected(store):
    pid = store.create_project("A", "codex")["id"]
    source = add_source(store, pid)
    add_source(store, pid, "corpus")
    runner = FakeRunner()
    rag = Rag(TinyEmbeddings())
    worker = Worker(store, runner, rag)
    queued = job(store, pid, source)
    await worker.execute(queued)
    assert store.rows("SELECT state FROM jobs WHERE id=?", (queued["id"],))[0]["state"] == "succeeded"
    assert runner.calls  # Missing index is repaired automatically.
    rag.build(store.workspace(pid), worker.corpus(pid))
    matches = rag.search_many(store.workspace(pid), worker.corpus(pid), ["医生"])[0]
    assert matches[0]["paragraph"] == 1
    assert matches[0]["score"] == 1.0
    queued = job(store, pid, source)
    await worker.execute(queued)
    assert store.rows("SELECT state FROM jobs WHERE id=?", (queued["id"],))[0]["state"] == "succeeded"
    assert '"references": [[' in runner.calls[0][2]
    assert '"source": "sample.docx"' in runner.calls[0][2]
    with pytest.raises(ValueError, match="重新构建"):
        rag.load(store.workspace(pid), [])


async def test_style_does_not_overwrite_concurrent_edit(store):
    pid = store.create_project("A", "codex")["id"]
    add_source(store, pid, "corpus")

    class EditingRunner(FakeRunner):
        async def run(self, *args, **kwargs):
            (store.workspace(pid) / "style.md").write_text("用户新编辑")
            return await super().run(*args, **kwargs)

    worker = Worker(store, EditingRunner(), Rag(TinyEmbeddings()))
    queued = store.enqueue(pid, "style", {})
    await worker.execute(queued)
    assert (store.workspace(pid) / "style.md").read_text() == "用户新编辑"
    assert store.rows("SELECT state FROM jobs")[0]["state"] == "needs_attention"


@pytest.mark.parametrize('other_agent', ['codex', 'codebuddy'])
async def test_project_parallel_queue_and_targeted_cancellation(store, other_agent):
    entered = {}
    release = asyncio.Event()
    active, peak = {}, {}

    class SlowRunner:
        async def run(self, pid, jid, *args, **kwargs):
            active[pid] = active.get(pid, 0) + 1
            peak[pid] = max(peak.get(pid, 0), active[pid])
            entered.setdefault(jid, asyncio.Event()).set()
            try:
                await release.wait()
                return "done"
            finally:
                active[pid] -= 1

    p1 = store.create_project("A", "codex")["id"]
    p2 = store.create_project("B", other_agent)["id"]
    worker = Worker(store, SlowRunner(), Rag(TinyEmbeddings()))
    first = store.enqueue(p1, "chat", {"message": "first"})
    following = store.enqueue(p1, "chat", {"message": "following"})
    second = store.enqueue(p2, "chat", {"message": "second"})
    task = asyncio.create_task(worker.loop())
    try:
        async def wait_started(jid):
            while jid not in entered:
                await asyncio.sleep(.01)
        await asyncio.wait_for(wait_started(first['id']), 2)
        await asyncio.wait_for(wait_started(second['id']), 2)
        assert following['id'] not in entered
        assert sum(active.values()) == 2
        assert not worker.cancel(p2, first['id'])
        assert worker.cancel(p1, first['id'])
        await asyncio.wait_for(wait_started(following['id']), 2)
        assert store.rows("SELECT state FROM jobs WHERE id=?", (second['id'],))[0]['state'] == 'running'
        release.set()
        for _ in range(100):
            await asyncio.sleep(0.01)
            if store.rows("SELECT state FROM jobs WHERE id=?", (second["id"],))[0]["state"] == "succeeded":
                break
        assert peak == {p1:1, p2:1}
        assert store.rows("SELECT state FROM jobs WHERE id=?", (first["id"],))[0]["state"] == "cancelled"
        assert store.rows("SELECT state FROM jobs WHERE id=?", (second["id"],))[0]["state"] == "succeeded"
        assert store.rows("SELECT state FROM jobs WHERE id=?", (following["id"],))[0]["state"] == "succeeded"
    finally:
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task


def test_restart_marks_running_interrupted_keeps_queued(store):
    pid = store.create_project("A", "codex")["id"]
    one = store.enqueue(pid, "chat", {"message": "1"})
    store.enqueue(pid, "chat", {"message": "2"})
    store.execute("UPDATE jobs SET state='running' WHERE id=?", (one["id"],))
    Worker(store).recover()
    assert [j["state"] for j in store.rows("SELECT * FROM jobs ORDER BY created")] == ["interrupted", "queued"]


def test_commands_pin_sessions_and_do_not_use_shell(tmp_path):
    codex = command("codex", "specific-session")
    assert codex[codex.index("resume") + 1] == "specific-session"
    assert "--last" not in codex
    assert "workspace-write" in codex
    buddy = command("codebuddy", "specific-session")
    assert buddy[buddy.index("--resume") + 1] == "specific-session"
    assert "--continue" not in buddy
    assert parse_json('```json\n{"passed":true}\n```') == {"passed": True}


def test_api_upload_config_conflict_and_download(tmp_path, monkeypatch):
    monkeypatch.setattr("transmux.app.availability", lambda: [{"id": "codex", "available": True}, {"id": "codebuddy", "available": True}])
    app = create_app(tmp_path, lambda store: Worker(store, FakeRunner(), Rag(TinyEmbeddings())))
    with TestClient(app) as client:
        assert client.get('/').status_code == 200
        assert client.get('/api/projects').json() == []
        project = client.post('/api/projects', json={"name": "工作区", "agent": "codex"})
        assert project.status_code == 201
        base = '/api/projects/' + project.json()["id"]
        original = client.get(base + '/config/style').json()
        changed = client.put(base + '/config/style', json={"content": "# Updated Style\n", "revision": original["revision"]})
        assert changed.status_code == 200
        conflict = client.put(base + '/config/style', json={"content": "旧写入", "revision": original["revision"]})
        assert conflict.status_code == 409
        assert client.put(base, json={"agent": "codebuddy"}).status_code == 405
        assert client.post('/api/projects', json={"name": "x", "agent": "invalid"}).status_code == 422
        assert client.post(base + '/files?kind=source', files={"file": ("bad.exe", b"x")}).status_code == 400
        assert client.post(base + '/files?kind=source', files={"file": ("bad.pdf", b"broken")}).status_code == 400
        uploaded = client.post(base + '/files?kind=source', files={"file": ("a.txt", b"Hello\nWorld")})
        assert uploaded.status_code == 201
        fid = uploaded.json()["id"]
        assert client.get(base + f'/files/{fid}/download').content == b"Hello\nWorld"
        assert client.post(base + '/jobs', json={"kind":"chat", "message":""}).status_code == 400
        assert client.post(base + '/jobs', json={"kind":"translate", "file_id":fid, "max_review_rounds":0}).status_code == 422
        assert client.post(base + '/jobs', json={"kind":"chat", "message":"x"}, headers={"Origin":"https://evil.example"}).status_code == 403
        assert client.get(base + '/artifact/runs/../style.md').status_code == 400
        queued = client.post(base + '/jobs', json={"kind":"translate", "file_id":fid})
        assert queued.status_code == 202
        import time
        for _ in range(100):
            result = client.get(base + '/jobs').json()[0]
            if result["state"] not in ("queued", "running"):
                break
            time.sleep(0.01)
        assert result["state"] == "succeeded"
        output_id = json.loads(result["result"])["file_id"]
        output = client.get(base + f'/files/{output_id}/download')
        assert output.status_code == 200
        assert output.content.startswith(b"PK")
        assert len(client.get(base + '/events').json()) >= 4


def test_export_preserves_tables_and_rejects_missing_blocks(tmp_path):
    source = tmp_path / 'original.docx'
    doc = Document()
    doc.add_paragraph('Title', style='Heading 1')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Cell'
    doc.save(source)
    with pytest.raises(ValueError, match='数量'):
        export_docx(source, ['标题'], tmp_path / 'out.docx')
    export_docx(source, ['标题', '单元格'], tmp_path / 'out.docx')
    assert extract(tmp_path / 'out.docx') == ['标题', '单元格']


async def test_chat_edits_cannot_change_approved_download(store):
    pid = store.create_project("A", "codex")["id"]
    source = add_source(store, pid)
    worker = Worker(store, FakeRunner(), Rag(TinyEmbeddings()))
    await worker.execute(job(store, pid, source))
    published = store.rows("SELECT * FROM files WHERE kind='output'")[0]
    original = store.download_path(pid, published).read_bytes()

    class OverwritingRunner:
        async def run(self, *args):
            doc = Document()
            doc.add_paragraph("修改后的文档")
            doc.save(store.safe_path(pid, published["path"]))
            return "已调整"

    worker.runner = OverwritingRunner()
    await worker.execute(store.enqueue(pid, "chat", {"message": "修改格式"}))
    edited = store.rows("SELECT * FROM files WHERE kind='edited'")[0]
    assert store.download_path(pid, published).read_bytes() == original
    assert extract(store.download_path(pid, edited)) == ["修改后的文档"]
