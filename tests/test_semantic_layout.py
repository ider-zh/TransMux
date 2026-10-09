import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
import pytest

from transmux.jobs import NeedsAttention
from transmux.presets import apply_preset, TEMPLATES
from transmux.semantic_layout import analyze, render, report, metadata_valid
from transmux.v2 import WorkspaceStore, WorkspaceWorker, create_v2_app
from test_layout import fake_pdf
from test_v2 import finished


def manuscript(path):
    doc = Document()
    doc.add_paragraph('Semantic layout study', 'Title')
    doc.add_paragraph('Abstract: We study reliable document transformations.')
    doc.add_paragraph('Keywords: rendering; citations')
    doc.add_paragraph('1 Introduction', 'Heading 1')
    paragraph = doc.add_paragraph('Prior work ')
    paragraph.add_run('[7]').bold = True
    paragraph.add_run(' is relevant; see Alpha (2020). ').italic = True
    paragraph.add_run('Unknown [99].')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '方法见 Beta study。'
    doc.add_paragraph('References', 'Heading 1')
    for text in ('[4] Alpha A. First study. Journal One, 2020, 1(2): 10-20.',
                 '[5] Beta B. Second study. Journal Two, 2021.',
                 '[7] Gamma G. Third study. Journal Three, 2022.',
                 '[8] Delta D. Uncited study. Journal Four, 2023.'):
        doc.add_paragraph(text)
    doc.save(path)


class SemanticRunner:
    bad_anchor = False

    async def run_isolated(self, pid, jid, prompt, schema=None):
        data = json.loads(prompt.rsplit('\n', 1)[-1])
        props = schema['properties']
        if 'blocks' in props:
            return {'blocks': [{'id': b['id'], 'role': b['role']} for b in data['blocks']]}
        if 'references' in props:
            # Unsupported metadata is explicitly retained, never fabricated.
            return {'references': []}
        if 'anchors' in props:
            anchors, unresolved = [], []
            for p in data['paragraphs']:
                for quote, rid in (('[7]', 'r3'), ('Alpha (2020)', 'r1'), ('Beta study', 'r2')):
                    if quote in p['text']:
                        anchors.append({'paragraph': p['id'], 'quote': 'invented' if self.bad_anchor else quote,
                                        'occurrence': 1, 'references': [rid]})
                if '[99]' in p['text']:
                    unresolved.append({'paragraph': p['id'], 'quote': '[99]', 'reason': 'No matching entry'})
            return {'anchors': anchors, 'unresolved': unresolved}
        raise AssertionError(props)


@pytest.mark.parametrize('template', ['jcst-submit', 'ieee-access'])
async def test_semantic_order_natural_anchors_uncited_and_object_preservation(tmp_path, template):
    store = WorkspaceStore(tmp_path / 'data')
    pid = store.create_project('Layout', 'codex')['id']
    job = store.enqueue(pid, 'layout', {'template': template})
    worker = WorkspaceWorker(store, SemanticRunner())
    source, rendered, output = (tmp_path / n for n in ('source.docx', 'semantic.docx', 'output.docx'))
    manuscript(source)
    before = source.read_bytes()
    plan = await analyze(worker, job, source, template)
    assert plan['order'] == ['r3', 'r1', 'r2', 'r4']
    assert plan['uncited'] == ['r4']
    roles = render(source, rendered, plan)
    apply_preset(rendered, output, template, roles, hashlib.sha256(rendered.read_bytes()).hexdigest())
    doc = Document(output)
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Prior work [1] is relevant; see [2].' in text
    assert '[?99]' in text
    assert doc.tables[0].cell(0, 0).text == '方法见 [3]。'
    refs = [p.text for p in doc.paragraphs if p.text.startswith('[')]
    assert len(refs) == 4
    assert refs[0].startswith('[1] Gamma') and refs[-1].startswith('[4] Delta')
    assert '未找到锚点的参考文献：1' in report(plan)
    body = next(p for p in doc.paragraphs if p.text.startswith('Prior work'))
    assert any(r.italic and 'is relevant' in r.text for r in body.runs)
    marker = next(r for r in body.runs if r.text == '[1]')
    assert marker.font.superscript == template.startswith('jcst')
    assert source.read_bytes() == before
    with ZipFile(source) as a, ZipFile(output) as b:
        for name in a.namelist():
            if name not in ('word/document.xml', 'word/styles.xml'):
                assert a.read(name) == b.read(name), name
    assert TEMPLATES[template]['citations']['retain_uncited'] is True


async def test_rejects_invented_anchors_and_stale_snapshots(tmp_path):
    store = WorkspaceStore(tmp_path / 'data')
    pid = store.create_project('Layout', 'codex')['id']
    job = store.enqueue(pid, 'layout', {})
    runner = SemanticRunner(); runner.bad_anchor = True
    source = tmp_path / 'source.docx'; manuscript(source)
    with pytest.raises(NeedsAttention, match='未出现在'):
        await analyze(WorkspaceWorker(store, runner), job, source, 'jcst-submit')
    with pytest.raises(NeedsAttention, match='原稿发生变化'):
        render(source, tmp_path / 'out.docx', {'source_sha256': 'wrong'})


def test_metadata_never_loses_or_invents_facts():
    from transmux.semantic_layout import FIELDS, reference_text
    meta = {k: '' for k in FIELDS}
    meta.update(type='article', authors='Alpha A', title='First study', container='Journal One', year='2020', volume='1', issue='2', pages='10-20')
    raw = 'Alpha A. First study. Journal One, 2020, 1(2): 10-20.'
    assert metadata_valid(raw, meta)
    ref = {'raw': raw, 'metadata': meta}
    ieee, applied = reference_text(ref, 'ieee')
    assert applied and 'vol. 1, no. 2, pp. 10-20, 2020' in ieee
    jcst, applied = reference_text(ref, 'jcst')
    assert applied and '2020, 1(2): 10-20' in jcst
    meta['doi'] = 'invented DOI'
    assert reference_text(ref, 'ieee') == (raw, False)
    meta['doi'] = ''
    ref['raw'] += ' Important additional bibliographic data.'
    assert reference_text(ref, 'ieee') == (ref['raw'], False)


async def test_repeated_ranges_non_citations_and_continuation_entries(tmp_path):
    store = WorkspaceStore(tmp_path / 'data')
    pid = store.create_project('Ranges', 'codex')['id']
    job = store.enqueue(pid, 'layout', {})
    source = tmp_path / 'source.docx'
    doc = Document()
    doc.add_paragraph('Reference processing', 'Title')
    doc.add_paragraph('First [3]–[4]. Again [3]. Vector [1,2].')
    doc.add_paragraph('References', 'Heading 1')
    doc.add_paragraph('[1] Uncited first entry.')
    doc.add_paragraph('[2] Uncited second entry.')
    doc.add_paragraph('[3] Cited third entry.')
    doc.add_paragraph('Continuation with DOI 10.1234/example.')
    doc.add_paragraph('[4] Cited fourth entry.')
    doc.save(source)

    class Runner(SemanticRunner):
        async def run_isolated(self, pid, jid, prompt, schema=None):
            data = json.loads(prompt.rsplit('\n', 1)[-1])
            if 'blocks' in schema['properties']:
                return {'blocks': [{'id': b['id'], 'role': b['role'],
                    'reference_continuation': b['text'].startswith('Continuation')} for b in data['blocks']]}
            if 'anchors' in schema['properties']:
                pid = next(p['id'] for p in data['paragraphs'] if 'Vector' in p['text'])
                return {'anchors': [], 'unresolved': [], 'non_citations': [{'paragraph': pid, 'quote': '[1,2]'}]}
            return await super().run_isolated(pid, jid, prompt, schema)

    plan = await analyze(WorkspaceWorker(store, Runner()), job, source, 'ieee-access')
    assert plan['order'] == ['r3', 'r4', 'r1', 'r2']
    assert plan['uncited'] == ['r1', 'r2']
    output = tmp_path / 'output.docx'
    render(source, output, plan)
    paragraphs = [p.text for p in Document(output).paragraphs]
    assert paragraphs[1] == 'First [1], [2]. Again [1]. Vector [1,2].'
    assert paragraphs[3:6] == ['[1] Cited third entry.', 'Continuation with DOI 10.1234/example.', '[2] Cited fourth entry.']
    assert paragraphs[-2:] == ['[3] Uncited first entry.', '[4] Uncited second entry.']


def test_semantic_api_publishes_docx_and_report_requires_explicit_input(tmp_path, monkeypatch):
    monkeypatch.setattr('transmux.app.availability', lambda: [{'id': 'codex', 'available': True}])
    monkeypatch.setattr('transmux.layout.render_pdf', fake_pdf)
    app = create_v2_app(tmp_path / 'data', lambda store: WorkspaceWorker(store, SemanticRunner()))
    source = tmp_path / 'source.docx'; manuscript(source)
    with TestClient(app) as client:
        pid = client.post('/api/projects', json={'name': 'Semantic', 'agent': 'codex'}).json()['id']
        base = f'/api/projects/{pid}'
        fid = client.post(base + '/attachments', files={'file': ('paper.docx', source.read_bytes())}).json()['id']
        assert client.post(base + '/messages', json={'kind': 'layout', 'template': 'jcst-submit'}).status_code == 400
        job = client.post(base + '/messages', json={'kind': 'layout', 'template': 'jcst-submit', 'file_ids': [fid]}).json()
        outcome = finished(client, pid, job['id'])
        assert outcome['state'] == 'needs_attention', outcome
        assert json.loads(outcome['result'])['status'] == 'draft'
        reviews = client.get(base + '/reviews').json()
        report_row = next(r for r in reviews if r['kind'] == 'layout_report')
        assert report_row['status'] == 'pending'
        content = client.get(base + f"/files/{report_row['file_id']}/download").text
        assert '未找到锚点的参考文献：1' in content
        assert any(r['kind'] == 'layout' and r['name'].endswith('.docx') for r in reviews)
        assert client.get(base + f'/files/{fid}/download').content == source.read_bytes()


@pytest.mark.parametrize('conflict', [False, True])
async def test_years_and_repeated_numeric_expressions_have_independent_locations(tmp_path, conflict):
    store = WorkspaceStore(tmp_path / 'data')
    pid = store.create_project('Position regression', 'codex')['id']
    job = store.enqueue(pid, 'layout', {})
    source = tmp_path / 'source.docx'
    doc = Document()
    doc.add_paragraph('Position regression', 'Title')
    doc.add_paragraph('His 1950 book is discussed (see Carnap-1950).')
    doc.add_paragraph('Published in 1956 (see McCarthy-1956).')
    doc.add_paragraph('Array [1] differs from citation [1].')
    doc.add_paragraph('References', 'Heading 1')
    doc.add_paragraph('[1] Carnap. Book. 1950.')
    doc.add_paragraph('[2] McCarthy. Paper. 1956.')
    doc.save(source)

    class Runner(SemanticRunner):
        async def run_isolated(self, pid, jid, prompt, schema=None):
            if 'anchors' not in schema['properties']:
                return await super().run_isolated(pid, jid, prompt, schema)
            assert 'occurrence' in schema['properties']['non_citations']['items']['required']
            result = {'anchors': [
                {'paragraph': 'p00002', 'quote': 'Carnap-1950', 'occurrence': 1, 'references': ['r1']},
                {'paragraph': 'p00003', 'quote': 'McCarthy-1956', 'occurrence': 1, 'references': ['r2']},
                {'paragraph': 'p00004', 'quote': '[1]', 'occurrence': 2, 'references': ['r1']}],
                'non_citations': [
                    {'paragraph': 'p00002', 'quote': '1950', 'occurrence': 1},
                    {'paragraph': 'p00003', 'quote': '1956', 'occurrence': 1},
                    {'paragraph': 'p00004', 'quote': '[1]', 'occurrence': 1}], 'unresolved': []}
            if conflict:
                result['anchors'].append({'paragraph': 'p00002', 'quote': '1950', 'occurrence': 2, 'references': ['r2']})
            return result

    if conflict:
        with pytest.raises(NeedsAttention, match='引文锚点重叠'):
            await analyze(WorkspaceWorker(store, Runner()), job, source, 'ieee-access')
        return
    plan = await analyze(WorkspaceWorker(store, Runner()), job, source, 'ieee-access')
    assert len(plan['anchors']) == 3 and plan['uncited'] == []
    output = tmp_path / 'output.docx'; render(source, output, plan)
    texts = [p.text for p in Document(output).paragraphs]
    assert texts[1] == 'His 1950 book is discussed [1].'
    assert texts[2] == 'Published in 1956 [2].'
    assert texts[3] == 'Array [1] differs from citation [1].'


def test_non_citation_legacy_and_invalid_positions_are_not_global_masks():
    from transmux.semantic_layout import non_citation_spans
    texts = {'p1': 'In 1950, see Carnap-1950. Array [1], citation [1].'}
    assert non_citation_spans([{'paragraph': 'p1', 'quote': '1950'}], texts) == {}
    item = {'paragraph': 'p1', 'quote': '[1]', 'occurrence': 1}
    assert non_citation_spans([item], texts) == {'p1': [(32, 35)]}
    for occurrence in (0, 3, True, None):
        with pytest.raises(NeedsAttention, match='出现位置'):
            non_citation_spans([dict(item, occurrence=occurrence)], texts)


@pytest.mark.parametrize('simple', [False, True])
def test_bibliography_label_and_external_anchor_preserve_word_field(simple):
    from docx.oxml import OxmlElement
    from transmux.semantic_layout import replace_span
    from transmux.presets import text_of
    doc = Document()
    p = doc.add_paragraph()
    if simple:
        field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'HYPERLINK "https://example.com"')
        run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'Author'
        run.append(text); field.append(run); p._p.append(field)
    else:
        for kind in ('begin', 'separate'):
            element = OxmlElement('w:fldChar'); element.set(qn('w:fldCharType'), kind)
            p.add_run()._r.append(element)
            if kind == 'begin':
                instruction = OxmlElement('w:instrText'); instruction.text = 'HYPERLINK "https://example.com"'
                p.add_run()._r.append(instruction)
        p.add_run('Author')
        element = OxmlElement('w:fldChar'); element.set(qn('w:fldCharType'), 'end'); p.add_run()._r.append(element)
    p.add_run(' see [7].')
    from lxml.etree import tostring
    field_xml = [tostring(child) for child in p._p][:-1]
    replace_span(p._p, 0, 0, '[1] ')
    assert text_of(p._p) == '[1] Author see [7].'
    assert [tostring(child) for child in p._p][1:-1] == field_xml
    start = text_of(p._p).index('[7]')
    replace_span(p._p, start, start + 3, '[2]')
    assert text_of(p._p) == '[1] Author see [2].'
    with pytest.raises(NeedsAttention, match='动态字段'):
        replace_span(p._p, 4, 10, 'Other')


async def test_paper_title_is_preserved_even_when_agent_calls_it_an_anchor(tmp_path):
    title = 'The Inversion of Functions Defined by Turing Machines'
    sentence = f'He contributed a short paper of just five pages to the collection, entitled “{title}” (see McCarthy-1956).'
    source = tmp_path / 'source.docx'
    doc = Document(); doc.add_paragraph('Title protection', 'Title'); doc.add_paragraph(sentence)
    doc.add_paragraph('References', 'Heading 1'); doc.add_paragraph(f'[1] McCarthy. {title}. 1956.'); doc.save(source)
    store = WorkspaceStore(tmp_path / 'data'); pid = store.create_project('Title protection', 'codex')['id']
    job = store.enqueue(pid, 'layout', {})
    class Runner(SemanticRunner):
        async def run_isolated(self, pid, jid, prompt, schema=None):
            if 'anchors' not in schema['properties']:
                return await super().run_isolated(pid, jid, prompt, schema)
            return {'anchors': [{'paragraph': 'p00002', 'quote': q, 'occurrence': 1, 'references': ['r1']}
                                for q in (title, 'McCarthy-1956')], 'unresolved': [], 'non_citations': []}
    plan = await analyze(WorkspaceWorker(store, Runner()), job, source, 'jcst-submit')
    assert [a['quote'] for a in plan['anchors']] == ['(see McCarthy-1956)']
    output = tmp_path / 'output.docx'; render(source, output, plan)
    assert Document(output).paragraphs[1].text == sentence.replace('(see McCarthy-1956)', '[1]')


@pytest.mark.parametrize('text,quote,expected', [
    ('Text (see [2]).', '[2]', 'Text [9].'),
    ('Text (See also Author-2020).', 'Author-2020', 'Text [9].'),
    ('正文（参见 [2]）。', '[2]', '正文[9]。'),
    ('Text (see [2] for details).', '[2]', 'Text (see [9] for details).'),
    ('Text (for details, see [2]).', '[2]', 'Text (for details, see [9]).'),
])
def test_see_only_parentheses(text, quote, expected):
    from transmux.semantic_layout import citation_wrapper_span
    start = text.index(quote)
    start, end = citation_wrapper_span(text, start, start + len(quote))
    assert text[:start] + '[9]' + text[end:] == expected
