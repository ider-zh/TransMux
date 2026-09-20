import json

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from transmux.alignment import source_blocks, normalize_groups
from transmux.documents import extract
from transmux.sections import ChapterPlanner
from transmux.jobs import Worker
from transmux.store import Store


def blocks(tmp_path, entries):
    source = tmp_path / 'chapters.docx'
    doc = Document()
    for text, level in entries:
        doc.add_heading(text, level) if level else doc.add_paragraph(text)
    doc.save(source)
    return source_blocks(source, extract(source))


def test_small_chapters_stay_whole_and_separate(tmp_path):
    data = blocks(tmp_path, [('Chapter A', 1), ('Intro', 0), ('Section A', 2), ('Body', 0),
                             ('Chapter B', 1), ('Other', 0)])
    planner = ChapterPlanner(data)
    first, carry = planner.window(0)
    assert len(first) == 4 and not carry
    assert [r['position'] for r in planner.window(4)[0]] == [5, 6]
    assert [n['title'] for n in planner.context(0, 4, [])['section_path']] == ['Chapter A']
    assert planner.context(0, 4, [])['read_only_context']['following_source'] is None
    # Chapter grouping never relaxes the existing paragraph-level heading boundary.
    with pytest.raises(ValueError, match='protected'):
        normalize_groups([dict(source_ids=[r['id'] for r in first], paragraphs=['Merged'], reason='Same chapter')], first)


def test_long_chapter_uses_subsections_without_heading_only_calls(tmp_path):
    data = blocks(tmp_path, [('C', 1), ('S1', 2), ('a' * 35, 0), ('S2', 2), ('b' * 35, 0), ('D', 1), ('End', 0)])
    planner = ChapterPlanner(data, max_chars=45)
    assert planner.units == [(0, 3), (3, 5), (5, 7)]
    assert len(planner.window(0)[0]) == 3
    assert len(planner.window(3)[0]) == 2
    context = planner.context(3, 5, [{'translation': 'Previous approved text.'}])
    assert [n['title'] for n in context['section_path']] == ['C', 'S2']
    assert context['read_only_context']['preceding_source']['source_id'] == 'p000003'
    assert context['read_only_context']['preceding_translation'] == 'Previous approved text.'
    assert context['read_only_context']['following_source'] is None


def test_nested_subsections_and_preamble_cover_every_source(tmp_path):
    data = blocks(tmp_path, [('Preamble', 0), ('C', 1), ('Intro' * 5, 0), ('S', 2),
                             ('Deep', 3), ('a' * 24, 0), ('Next', 3), ('b' * 24, 0), ('End', 1)])
    planner = ChapterPlanner(data, max_chars=35)
    assert [i for start, end in planner.units for i in range(start, end)] == list(range(len(data)))
    assert planner.units[0] == (0, 1)
    assert (3, 6) in planner.units
    assert (6, 8) in planner.units
    assert [n['title'] for n in planner.context(5, 6, [])['section_path']] == ['C', 'S', 'Deep']


def test_plain_text_conservative_headings_and_fenced_code(tmp_path):
    source = tmp_path / 'source.md'
    source.write_text('# Intro\nbody\n## Methods\n1. First item\nChapter 2 describes the methods.\n```\n# not a heading\n```\nChapter 2: Results\n正文\n第一章 绪论\n文本\n')
    data = source_blocks(source, extract(source))
    assert [(i, b['level']) for i, b in enumerate(data) if b['kind'] == 'heading'] == [(0, 1), (2, 2), (8, 1), (10, 1)]
    assert all(data[i]['protected'] for i in [3, 5, 6, 7])


def test_docx_inherited_outline_heading_and_table_is_not_chapter(tmp_path):
    source = tmp_path / 'source.docx'
    doc = Document()
    custom = doc.styles.add_style('Custom Section', WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = doc.styles['Heading 2']
    doc.add_paragraph('Inherited', custom)
    p = doc.add_paragraph('Explicit')
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), '2')
    p._p.get_or_add_pPr().append(outline)
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = 'Cell title'
    cell.paragraphs[0].style = 'Heading 1'
    doc.save(source)
    data = source_blocks(source, extract(source))
    assert [b['kind'] for b in data] == ['heading', 'heading', 'table']
    assert [b['level'] for b in data[:2]] == [2, 3]


def test_length_budget_carry_and_single_large_paragraph_progress(tmp_path):
    data = blocks(tmp_path, [('C', 1)] + [('a' * 10, 0)] * 7 + [('D', 1), ('end', 0)])
    planner = ChapterPlanner(data, max_chars=35)
    first, carry = planner.window(0)
    assert len(first) == 4 and carry
    # Revisit the previous window's provisional tail; never carry into the next chapter.
    offset = len(first) - 1
    assert planner.window(offset)[0][0]['id'] == first[-1]['id']
    assert not planner.window(6)[1]
    large = blocks(tmp_path, [('x' * 80, 0), ('y' * 80, 0)])
    planner = ChapterPlanner(large, max_chars=40)
    assert len(planner.window(0)[0]) == 1 and not planner.window(0)[1]
    assert planner.context(0, 1, [])['oversized_paragraph']


def test_token_estimate_and_context_are_bounded(tmp_path):
    data = blocks(tmp_path, [('中' * 100, 0)] * 4)
    planner = ChapterPlanner(data, max_chars=10000, max_estimated_tokens=800)
    assert len(planner.window(0)[0]) == 2
    data = blocks(tmp_path, [('a' * 2000, 0), ('b', 0), ('c' * 2000, 0)])
    context = ChapterPlanner(data).context(1, 2, [{'translation': 'x' * 2000}])['read_only_context']
    assert len(context['preceding_source']['text']) == len(context['following_source']['text']) == 800
    assert context['preceding_source']['truncated']
    assert len(context['preceding_translation']) == 800


async def test_worker_injects_read_only_chapter_context_and_progress(tmp_path):
    store = Store(tmp_path / 'data')
    pid = store.create_project('sections', 'codex')['id']
    work = store.workspace(pid)
    source = work / 'sources' / 'source.md'
    source.write_text('# A\nFirst body.\n## Nested\nMore body.\n# B\nOther body.')
    source.with_suffix('.md.json').write_text(json.dumps(extract(source)))
    fid = store.add_file(pid, source.name, 'source', source)

    class Runner:
        batches = []

        async def run(self, pid, jid, prompt, schema):
            if 'translations' not in schema['properties']:
                return dict(passed=True, issues=[], mapping_issues=[])
            files = sorted((work / 'runs' / jid).glob('batch-*-input.json'))
            request = json.loads(files[-1].read_text())
            self.batches.append(request)
            assert 'not additional translation input' in prompt
            return dict(translations=[dict(source_ids=[b['id']], paragraphs=[t], reason='')
                                      for b, t in zip(request['source_blocks'], request['source'])], mappings=[])

    runner = Runner()
    task = store.enqueue(pid, 'translate', dict(file_id=fid, use_rag=False, max_review_rounds=1))
    await Worker(store, runner).execute(task)
    finished = store.rows('SELECT state,result FROM jobs WHERE id=?', (task['id'],))[0]
    assert finished['state'] == 'succeeded', finished
    assert [len(b['source']) for b in runner.batches] == [4, 2]
    assert runner.batches[1]['section_path'][0]['title'] == '# B'
    assert runner.batches[1]['read_only_context']['preceding_source'] is None
    events = store.rows('SELECT text FROM events WHERE job=?', (task['id'],))
    assert any('# B' in e['text'] and '5–6' in e['text'] for e in events)
    store.db.close()
