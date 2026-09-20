import hashlib
import io
import json
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
import pytest

from transmux.app import create_app
from transmux.presets import ACCESS, JCST, apply_access, apply_preset, content_signature, read_package, structure
from test_comparison import wait_job
from test_layout import fake_pdf, worker


def article(path):
    doc = Document()
    doc.add_paragraph('An Article on Reliable Export', 'Title')
    doc.add_paragraph('A. Example and B. Sample')
    doc.add_paragraph('Department of Computing, Example University')
    doc.add_paragraph('ABSTRACT We evaluate document preservation during formatting.')
    doc.add_paragraph('INDEX TERMS Document processing, preservation, formatting.')
    doc.add_paragraph('I. INTRODUCTION', 'Heading 1')
    paragraph = doc.add_paragraph('Existing citation [1]. ')
    paragraph.add_run('Important emphasis.').italic = True
    paragraph.add_run('Bold text.').bold = True
    doc.add_paragraph('First numbered item', 'List Number')
    doc.add_paragraph('Second numbered item', 'List Number')
    eq = OxmlElement('m:oMath')
    mr = OxmlElement('m:r')
    mt = OxmlElement('m:t')
    mt.text = 'x=1'
    mr.append(mt)
    eq.append(mr)
    paragraph._p.append(eq)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Merged cell'
    table.cell(1, 0).text = 'Original data'
    doc.add_paragraph('REFERENCES', 'Heading 1')
    doc.add_paragraph('[1] A. Example, Original title, 2024.')
    doc.sections[0].header.paragraphs[0].text = 'Original running header'
    doc.save(path)


def test_access_preserves_content_assets_and_applies_verified_geometry(tmp_path):
    source, output = tmp_path/'input.docx', tmp_path/'output.docx'
    article(source)
    before = source.read_bytes()
    info = structure(source)
    roles = {'b00002':'authors', 'b00003':'affiliation'}
    result = apply_access(source, output, roles, info['source_sha256'])
    assert source.read_bytes() == before
    assert output.read_bytes() != before
    original, _ = read_package(source)
    formatted, _ = read_package(output)
    assert content_signature(original) == content_signature(formatted)
    with ZipFile(source) as a, ZipFile(output) as b:
        assert a.namelist() == b.namelist()
        for part in a.namelist():
            if part not in ('word/document.xml', 'word/styles.xml'):
                assert a.read(part) == b.read(part), part
    doc = Document(output)
    assert len(doc.sections) == 2
    assert doc.sections[0]._sectPr.find(qn('w:cols')).get(qn('w:num'), '1') == '1'
    assert doc.sections[1]._sectPr.find(qn('w:cols')).get(qn('w:num')) == '2'
    assert doc.sections[1]._sectPr.find(qn('w:cols')).get(qn('w:space')) == '400'
    assert doc.sections[1].page_width.twips == 11520
    assert doc.sections[1].page_height.twips == 15660
    assert doc.sections[1].left_margin.twips == 740
    assert doc.paragraphs[0].style.font.size.pt == 22
    assert doc.paragraphs[0].style.font.name == 'Helvetica'
    body = next(p for p in doc.paragraphs if p.text.startswith('Existing citation'))
    assert body.style.font.size.pt == 10
    assert any(r.italic for r in body.runs) and any(r.bold for r in body.runs)
    assert doc.tables[0].cell(1, 0).text == 'Original data'
    for p in doc.paragraphs:
        if 'numbered item' in p.text:
            assert p._p.pPr.numPr.numId.val == Document(source).styles['List Number'].element.pPr.numPr.numId.val
    assert any(f['code'] == 'complex_block' and f['block_id'] for f in result['findings'])
    assert not any(f['code'] == 'missing_authors' for f in result['findings'])
    assert any(f['code'] == 'references_pending' for f in result['findings'])
    assert ACCESS['source_sha256'] and ACCESS['source_url'].endswith('Access-Template-2024.docx')


def test_structure_overrides_stale_validation_and_existing_section_preservation(tmp_path):
    source, output = tmp_path/'input.docx', tmp_path/'output.docx'
    article(source)
    info = structure(source)
    assert info['blocks'][1]['role'] == 'frontmatter'
    with pytest.raises(ValueError, match='文稿已变化'):
        apply_access(source, output, {}, '0'*64)
    with pytest.raises(ValueError, match='无效段落'):
        apply_access(source, output, {'b99999':'body'})
    doc = Document(source)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph('An appendix paragraph.')
    doc.save(source)
    result = apply_access(source, output)
    doc = Document(output)
    assert len(doc.sections) == 3
    assert doc.sections[-1].start_type == WD_SECTION_START.NEW_PAGE
    assert any(f['code'] == 'existing_sections' for f in result['findings'])


@pytest.mark.parametrize('template,metadata', [('ieee-access', ACCESS), ('jcst', JCST)])
def test_preset_export_api_draft_history_no_model_call_and_original_unchanged(tmp_path, monkeypatch, template, metadata):
    monkeypatch.setattr('transmux.app.availability', lambda: [{'id':'codex','available':True}])
    monkeypatch.setattr('transmux.layout.render_pdf', fake_pdf)
    source = tmp_path/'input.docx'
    article(source)
    original = source.read_bytes()
    app = create_app(tmp_path/'app', worker)
    with TestClient(app) as client:
        pid = client.post('/api/projects', json={'name':'layout','agent':'codex'}).json()['id']
        other = client.post('/api/projects', json={'name':'other','agent':'codex'}).json()['id']
        base = '/api/projects/'+pid
        fid = client.post(base+'/files?kind=manuscript', files={'file':('paper.docx', original)}).json()['id']
        info = client.get(base+f'/files/{fid}/structure').json()
        assert client.get(f'/api/projects/{other}/files/{fid}/structure').status_code == 400
        capabilities = client.get('/api/layout/capabilities').json()
        assert metadata in capabilities['templates']
        assert client.post(base+'/jobs', json={'kind':'layout','file_id':fid,'template':template}).status_code == 400
        payload = {'kind':'layout','file_id':fid,'template':template,'source_sha256':info['source_sha256'],
                   'roles':{'b00002':'authors','b00003':'affiliation'}}
        job = client.post(base+'/jobs', json=payload).json()
        result = wait_job(client, base, job['id'])
        assert result['state'] == 'needs_attention', result
        assert json.loads(result['result'])['status'] == 'draft'
        row = client.get(base+f'/files/{fid}/exports').json()[0]
        manifest = row['manifest']
        assert manifest['template'] == metadata
        assert manifest['source_sha256'] == hashlib.sha256(original).hexdigest()
        assert manifest['docx_sha256'] != manifest['source_sha256']
        assert manifest['pdf'] and manifest['findings'] and manifest['status'] == 'draft'
        download = client.get(base+f'/exports/{row["id"]}/docx').content
        assert len(Document(io.BytesIO(download)).sections) == 2
        assert client.get(base+f'/files/{fid}/download').content == original
        assert app.state.worker.runner.calls == []
        # Original format is still the independent default, even after applying a preset.
        task = client.post(base+'/jobs', json={'kind':'layout','file_id':fid}).json()
        result = wait_job(client, base, task['id'])
        assert result['state'] == 'succeeded'
        rows = client.get(base+f'/files/{fid}/exports').json()
        assert len(rows) == 2 and rows[0]['manifest']['template']['id'] == 'original'
        assert client.get(base+f'/exports/{rows[0]["id"]}/docx').content == original
        assert client.post(base+'/jobs', json={**payload,'roles':{'b00002':'invalid'}}).status_code == 422


def test_jcst_geometry_typography_and_content_preservation(tmp_path):
    source, output = tmp_path/'input.docx', tmp_path/'output.docx'
    article(source)
    doc = Document(source)
    doc.add_paragraph('2.1 Evaluation', 'Heading 2')
    doc.add_paragraph('2.1.1 Details', 'Heading 3')
    doc.add_paragraph('Fig. 1. A preservation example', 'Caption')
    doc.save(source)
    original = source.read_bytes()
    info = structure(source)
    result = apply_preset(source, output, 'jcst', {'b00002':'authors','b00003':'affiliation'}, info['source_sha256'])
    assert source.read_bytes() == original
    assert content_signature(read_package(source)[0]) == content_signature(read_package(output)[0])
    with ZipFile(source) as a, ZipFile(output) as b:
        assert a.namelist() == b.namelist()
        for part in a.namelist():
            if part not in ('word/document.xml', 'word/styles.xml'):
                assert a.read(part) == b.read(part), part
    doc = Document(output)
    assert len(doc.sections) == 2
    first, body_section = doc.sections
    assert first._sectPr.find(qn('w:cols')).get(qn('w:num'), '1') == '1'
    assert body_section._sectPr.find(qn('w:cols')).get(qn('w:num')) == '2'
    assert body_section._sectPr.find(qn('w:cols')).get(qn('w:space')) == '421'
    assert (body_section.page_width.twips, body_section.page_height.twips) == (11907,16840)
    assert (body_section.top_margin.twips, body_section.bottom_margin.twips) == (1474,765)
    assert body_section.left_margin.twips == body_section.right_margin.twips == 839
    assert doc.paragraphs[0].style.font.name == 'Times New Roman'
    assert doc.paragraphs[0].style.font.size.pt == 16
    assert doc.paragraphs[0].style.font.bold
    assert doc.paragraphs[2].style.font.italic
    body = next(p for p in doc.paragraphs if p.text.startswith('Existing citation'))
    assert body.style.font.size.pt == 12
    assert body.style.paragraph_format.line_spacing == 1.5
    assert any(r.italic for r in body.runs) and any(r.bold for r in body.runs)
    heading = next(p for p in doc.paragraphs if p.text.startswith('2.1.1'))
    assert heading.style.font.italic and heading._p.pPr.outlineLvl.val == 2
    caption = next(p for p in doc.paragraphs if p.text.startswith('Fig. 1'))
    assert caption.style.font.size.pt == 10
    assert any(f['code'] == 'references_pending' for f in result['findings'])
    assert 'IEEE' not in json.dumps(result, ensure_ascii=False)
    assert JCST['version'] == '2022.06-r1' and JCST['source_sha256']
    with pytest.raises(ValueError, match='文稿已变化'):
        apply_preset(source, output, 'jcst', {}, '0'*64)
