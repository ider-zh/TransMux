"""Long-document layout check with real PDF: PYTHONPATH=. python tests/browser_layout.py."""
import io
from pathlib import Path
import socket
import tempfile
import threading
import time

import httpx
import uvicorn
from docx import Document
from playwright.sync_api import sync_playwright

from transmux.app import create_app
from test_layout import worker


def main():
    doc = Document()
    doc.add_paragraph('Reliable Academic Document Export', 'Title')
    doc.add_paragraph('A. Example and B. Sample')
    doc.add_paragraph('Department of Computing, Example University')
    doc.add_paragraph('Abstract This sample checks document layout and long-list navigation.')
    doc.add_paragraph('Keywords document processing, formatting, verification')
    doc.add_paragraph('1. Introduction', 'Heading 1')
    for index in range(45):
        doc.add_paragraph(f'Paragraph {index+1}. ' + 'Each export preserves the text while applying the selected journal layout. ' * 4)
    doc.add_paragraph('References', 'Heading 1')
    doc.add_paragraph('[1] Example A. Document processing. 2024.')
    raw = io.BytesIO()
    doc.save(raw)
    # Snap LibreOffice has a private /tmp; use the same workspace filesystem as the service.
    with tempfile.TemporaryDirectory(prefix='transmux-layout-browser-', dir=Path(__file__).resolve().parents[1]) as root:
        app = create_app(root, worker)
        listener = socket.socket()
        listener.bind(('127.0.0.1', 0))
        address = f'http://127.0.0.1:{listener.getsockname()[1]}'
        server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
        thread = threading.Thread(target=server.run, kwargs={'sockets':[listener]})
        thread.start()
        try:
            with httpx.Client(base_url=address, trust_env=False) as client:
                for _ in range(100):
                    try:
                        if client.get('/api/health').status_code == 200:
                            break
                    except httpx.ConnectError:
                        pass
                    time.sleep(.1)
                pid = client.post('/api/projects', json={'name':'Layout verification','agent':'codex'}).json()['id']
                base = '/api/projects/'+pid
                fid = client.post(base+'/files?kind=manuscript', files={'file':('long-article.docx',raw.getvalue())}).json()['id']
                with sync_playwright() as playwright:
                    browser = playwright.chromium.launch(executable_path='/usr/bin/google-chrome', headless=True, args=['--no-sandbox'])
                    page = browser.new_page(viewport={'width':1366,'height':768})
                    errors = []
                    page.on('pageerror', lambda error: errors.append(str(error)))
                    page.goto(address)
                    page.locator(f'[data-layout="{fid}"]').click()
                    page.locator('#layout-template').select_option('jcst')
                    page.locator('#layout-tab-structure').click()
                    page.locator('[data-layout-role="b00002"]').select_option('authors')
                    page.locator('[data-layout-role="b00003"]').select_option('affiliation')
                    panel = page.locator('#layout-structure-list')
                    assert panel.evaluate('el => el.scrollHeight > el.clientHeight * 3')
                    page.locator('[data-layout-role]').last.scroll_into_view_if_needed()
                    page.locator('[data-layout-role]').last.select_option('reference')
                    scroll = panel.evaluate('el => el.scrollTop')
                    assert scroll > 0
                    page.locator('#layout-tab-preview').click()
                    page.locator('#layout-tab-structure').click()
                    assert abs(panel.evaluate('el => el.scrollTop')-scroll) <= 1
                    page.locator('#layout-generate').click()
                    page.locator('#layout-pdf').wait_for(state='visible', timeout=120000)
                    page.wait_for_function("!document.querySelector('#layout-generate').disabled")
                    pdf_box = page.locator('#layout-pdf').bounding_box()
                    assert pdf_box['width'] == 1366 and pdf_box['height'] > 450, pdf_box
                    row = client.get(base+f'/files/{fid}/exports').json()[0]
                    assert row['manifest']['renderer']['pages'] >= 3
                    assert client.get(base+f'/files/{fid}/download').content == raw.getvalue()
                    # Native PDF viewer renders asynchronously; give it time for the visual capture.
                    page.wait_for_timeout(1500)
                    page.screenshot(path='/tmp/transmux-layout-real-pdf.png')
                    for width,height in [(1920,1080),(390,844)]:
                        page.set_viewport_size({'width':width,'height':height})
                        assert page.locator('#layout-pdf').bounding_box()['width'] == width
                        assert page.locator('#layout-dialog').evaluate('el => el.scrollWidth === el.clientWidth')
                        page.locator('#layout-tab-structure').click()
                        page.locator('[data-layout-role]').last.scroll_into_view_if_needed()
                        page.locator('[data-layout-role]').last.select_option('reference')
                        assert panel.evaluate('el => el.scrollTop > 0')
                        page.locator('#layout-tab-preview').click()
                    assert not errors, errors
                    assert app.state.worker.runner.calls == []
                    browser.close()
                    print('Long-document browser check passed: real multi-page JCST PDF, full-width preview, structure scrolling and tab state at desktop/mobile sizes; no Agent calls.')
        finally:
            server.should_exit = True
            thread.join(timeout=10)
            listener.close()


if __name__ == '__main__':
    main()
