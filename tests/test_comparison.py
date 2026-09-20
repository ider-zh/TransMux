import json
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient
from lxml import etree

from transmux.app import create_app
from transmux.documents import doc_paragraphs, export_docx, extract, legacy_alignment, revise_docx
from transmux.jobs import Worker
from transmux.rag import Rag
from transmux.store import ConfigConflict, Store
from test_workflows import FakeRunner, TinyEmbeddings, add_source, job


class RevisionRunner(FakeRunner):
    async def run(self, pid, jid, prompt, schema=None):
        if schema and 'translation' in schema['properties']:
            self.calls.append((pid, jid, prompt, schema))
            return {'translation':'医生在医院提供医疗服务。'}
        return await super().run(pid, jid, prompt, schema)


async def setup_translation(store):
    pid = store.create_project('test', 'codex', target_language='zh-CN')['id']
    source = add_source(store, pid)
    runner = RevisionRunner()
    worker = Worker(store, runner, Rag(TinyEmbeddings()))
    task = job(store, pid, source)
    await worker.execute(task)
    result = store.rows('SELECT * FROM jobs WHERE id=?', (task['id'],))[0]
    assert result['state'] == 'succeeded'
    return pid, source, json.loads(result['result'])['file_id'], worker


async def test_comparison_and_revision_preserve_original_and_other_paragraphs(tmp_path):
    store = Store(tmp_path)
    pid, source, fid, worker = await setup_translation(store)
    comparison = store.comparison(pid, fid)
    assert comparison['source_file'] == source and comparison['root_file'] == fid
    assert comparison['paragraphs'][0]['kind'] == 'heading'
    assert comparison['paragraphs'][1]['kind'] == 'table'
    assert comparison['paragraphs'][1]['row'] == comparison['paragraphs'][1]['column'] == 1
    published = store.download_path(pid, store.file(pid, fid))
    original_bytes = published.read_bytes()
    old_xml = [etree.tostring(p._p) for p in doc_paragraphs(Document(published))]
    worker.runner.calls.clear()
    revision = store.enqueue(pid, 'revise', dict(file_id=fid, paragraph=1, message='更简洁', use_rag=False, max_review_rounds=2))
    await worker.execute(revision)
    result = store.rows('SELECT * FROM jobs WHERE id=?', (revision['id'],))[0]
    assert result['state'] == 'succeeded'
    assert len(worker.runner.calls) == 2
    revised_id = json.loads(result['result'])['file_id']
    new = store.comparison(pid, revised_id)
    assert new['parent_file'] == fid and new['root_file'] == fid
    assert new['paragraphs'][0]['original'] == comparison['paragraphs'][0]['original']
    assert new['paragraphs'][0]['translation'] == '医生在医院提供医疗服务。'
    assert new['paragraphs'][1:] == comparison['paragraphs'][1:]
    revised_path = store.download_path(pid, store.file(pid, revised_id))
    assert extract(revised_path) == [row['translation'] for row in new['paragraphs']]
    assert published.read_bytes() == original_bytes
    assert [etree.tostring(p._p) for p in doc_paragraphs(Document(revised_path))][1:] == old_xml[1:]
    with pytest.raises(ConfigConflict, match='新'):
        store.require_revision_available(pid, fid)
    # Further changes must build on the latest version, preserving the earlier revision.
    second = store.enqueue(pid, 'revise', dict(file_id=revised_id, paragraph=2, message='调整措辞', use_rag=False, max_review_rounds=1))
    await worker.execute(second)
    result = json.loads(store.rows('SELECT result FROM jobs WHERE id=?', (second['id'],))[0]['result'])
    assert store.comparison(pid, result['file_id'])['paragraphs'][0] == new['paragraphs'][0]
    store.delete_project(pid, 'test')
    assert store.rows('SELECT * FROM comparisons') == []
    store.db.close()


async def test_failed_revision_keeps_old_version_and_saves_drafts(tmp_path):
    store = Store(tmp_path)
    pid, _, fid, worker = await setup_translation(store)
    old = store.comparison(pid, fid)
    worker.runner.passed = False
    worker.runner.calls.clear()
    revision = store.enqueue(pid, 'revise', dict(file_id=fid, paragraph=2, message='调整', use_rag=False, max_review_rounds=2))
    with pytest.raises(ConfigConflict, match='已有段落修改'):
        store.require_revision_available(pid, fid)
    await worker.execute(revision)
    assert store.rows('SELECT state FROM jobs WHERE id=?', (revision['id'],))[0]['state'] == 'needs_attention'
    assert len(worker.runner.calls) == 4
    assert store.comparison(pid, fid) == old
    assert len(store.rows('SELECT * FROM comparisons')) == 1
    assert (store.workspace(pid) / 'runs' / revision['id'] / 'revision-draft-2.json').is_file()
    assert len(store.rows("SELECT * FROM files WHERE kind='output'")) == 1
    store.db.close()


def wait_job(client, base, jid):
    for _ in range(200):
        task = next(job for job in client.get(base + '/jobs').json() if job['id'] == jid)
        if task['state'] not in ('running','queued'):
            return task
        time.sleep(.02)
    raise AssertionError('Job did not finish')


def test_comparison_api_legacy_recovery_bounds_isolation_and_latest_version(tmp_path, monkeypatch):
    monkeypatch.setattr('transmux.app.availability', lambda: [{'id':'codex','available':True}])
    app = create_app(tmp_path, lambda store: Worker(store, RevisionRunner(), Rag(TinyEmbeddings())))
    with TestClient(app) as client:
        pid = client.post('/api/projects', json=dict(name='test', agent='codex', target_language='zh-CN')).json()['id']
        other = client.post('/api/projects', json=dict(name='other', agent='codex')).json()['id']
        base = f'/api/projects/{pid}'
        source = client.post(base + '/files?kind=source', files={'file':('sample.txt',b'The doctor works in a hospital.\nKeep the second paragraph.')}).json()['id']
        task = client.post(base + '/jobs', json=dict(kind='translate', file_id=source, use_rag=False)).json()
        result = wait_job(client, base, task['id'])
        assert result['state'] == 'succeeded'
        fid = json.loads(result['result'])['file_id']
        url = base + f'/files/{fid}/comparison'
        old = client.get(url).json()
        assert old['paragraphs'][0]['original'] == 'The doctor works in a hospital.'
        assert client.get(f'/api/projects/{other}/files/{fid}/comparison').status_code == 400
        assert client.get(base + f'/files/{source}/comparison').status_code == 400
        assert client.post(base + '/jobs', json=dict(kind='revise', file_id=fid, paragraph=99, message='x')).status_code == 400
        assert client.post(base + '/jobs', json=dict(kind='revise', file_id=fid, paragraph=1)).status_code == 400
        # Simulate an older approved output without an alignment record.
        import sqlite3
        with sqlite3.connect(tmp_path / 'transmux.sqlite3') as db:
            db.execute('DELETE FROM comparisons WHERE file=?', (fid,))
            source_path = db.execute('SELECT path FROM files WHERE id=?', (source,)).fetchone()[0]
        (tmp_path / 'projects' / pid / source_path).write_text('Changed source text outside the application')
        assert next(f for f in client.get(base + '/files').json() if f['id'] == fid)['can_compare']
        recovered = client.get(url)
        assert recovered.status_code == 200
        assert recovered.json()['paragraphs'] == old['paragraphs']
        task = client.post(base + '/jobs', json=dict(kind='revise', file_id=fid, paragraph=1, message='更准确', use_rag=False)).json()
        revised = wait_job(client, base, task['id'])
        assert revised['state'] == 'succeeded'
        new_id = json.loads(revised['result'])['file_id']
        data = client.get(base + f'/files/{new_id}/comparison').json()
        assert len(data['versions']) == 2 and data['parent_file'] == fid
        assert client.post(base + '/jobs', json=dict(kind='revise', file_id=fid, paragraph=1, message='x')).status_code == 409
    with TestClient(app) as client:
        assert client.get(base + f'/files/{new_id}/comparison').json()['paragraphs'] == data['paragraphs']


async def test_revision_detects_changed_published_content(tmp_path):
    store = Store(tmp_path)
    pid, _, fid, worker = await setup_translation(store)
    published = store.download_path(pid, store.file(pid, fid))
    doc = Document(published)
    doc.paragraphs[0].text = 'unexpected external edit'
    doc.save(published)
    task = store.enqueue(pid, 'revise', dict(file_id=fid, paragraph=1, message='调整', use_rag=False, max_review_rounds=1))
    await worker.execute(task)
    assert store.rows('SELECT state FROM jobs WHERE id=?', (task['id'],))[0]['state'] == 'failed'
    assert len(store.rows('SELECT * FROM comparisons')) == 1
    store.db.close()


def test_textual_controls_do_not_break_alignment_and_legacy_requires_snapshots(tmp_path):
    source = tmp_path / 'source.docx'
    doc = Document()
    doc.add_paragraph('Original\ttext\nwith a line break.')
    doc.add_paragraph('Second paragraph.')
    doc.save(source)
    output = tmp_path / 'output.docx'
    export_docx(source, ['Translated first paragraph.', 'Translated second paragraph.'], output)
    assert extract(output) == ['Translated first paragraph.', 'Translated second paragraph.']
    revised = tmp_path / 'revised.docx'
    revise_docx(output, extract(output), 1, 'Revised second paragraph.', revised)
    assert extract(revised) == ['Translated first paragraph.', 'Revised second paragraph.']
    with pytest.raises(ValueError, match='缺少可靠'):
        legacy_alignment(tmp_path, 'missing-job', source, output)
