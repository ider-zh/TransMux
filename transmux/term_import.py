"""Table import is parsed and previewed before any project data is changed."""
import csv
import io
from zipfile import BadZipFile, ZipFile

from . import terminology as terms


def parse_table(data, filename='paste.tsv'):
    if len(data) > 2 * 1024 * 1024:
        raise ValueError('导入文件超过 2 MB，请拆分')
    if filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            with ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError('Excel 解压后超过 20 MB')
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        except (BadZipFile, KeyError) as exc:
            raise ValueError('无法读取 XLSX 文件，请重新导出有效的 Excel 工作簿') from exc
        try:
            sheet = workbook.worksheets[0]
            if sheet.max_column > 30 or sheet.max_row > 5001:
                raise ValueError('最多 5000 条、30 列，请拆分表格')
            values = []
            for row in sheet.iter_rows():
                if len(row) > 30 or len(values) >= 5001:
                    raise ValueError('最多 5000 条、30 列，请拆分表格')
                if any(c.data_type == 'f' for c in row):
                    raise ValueError('请先将 Excel 公式转为值再导入')
                values.append(['' if c.value is None else str(c.value) for c in row])
            sheet_name = sheet.title
        finally:
            workbook.close()
    else:
        if not filename.lower().endswith(('.csv', '.tsv', '.txt')):
            raise ValueError('支持 XLSX、UTF-8 CSV/TSV 或粘贴表格')
        try:
            text = data.decode('utf-8-sig')
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',\t;')
            except csv.Error:
                dialect = csv.excel_tab
            values = list(csv.reader(io.StringIO(text), dialect))
        except (UnicodeError, csv.Error) as exc:
            raise ValueError('请使用 UTF-8 CSV，或从 Excel 直接复制粘贴') from exc
        sheet_name = ''
    values = [row for row in values if any(c.strip() for c in row)]
    if not values or len(values) > 5001 or any(len(row) > 30 or any(len(c) > 10000 for c in row) for row in values):
        raise ValueError('表格为空或超过 5000 条、30 列、单格 10000 字符的限制')
    return {'cells': values, 'sheet': sheet_name}


def candidate(kind, values, source='表格导入'):
    if kind not in terms.KINDS:
        raise ValueError('无效术语类型')
    allowed = {*terms.FIELDS[kind], 'scope', 'reason', 'evidence'}
    if not isinstance(values, dict) or set(values) - allowed or any(not isinstance(v, str) for v in values.values()):
        raise ValueError('无效导入字段')
    row = {field: values.get(field, '').strip() for field in terms.FIELDS[kind]}
    row.update({field: values[field].strip() for field in ('scope', 'reason', 'evidence') if field in values})
    row.update(source=row.get('source') or source, origin='manual', status='active')
    if kind == 'people' and not row['translation']:
        row['status'] = 'pending'
    terms.decode(kind, terms.encode(dict(rows=[row], deleted=[], legacy='')))
    return row


def preview(kind, document, values):
    if not isinstance(values, list) or len(values) > 5000:
        raise ValueError('最多导入 5000 条')
    known = {terms.key(kind, r): r for r in document['rows'] if not r.get('variant')}
    output = []
    for value in values:
        try:
            row = candidate(kind, value)
            identity = terms.key(kind, row)
            old = known.get(identity)
            compare = [*terms.FIELDS[kind], 'scope']
            duplicate = old is not None and all(old.get(f, '') == row.get(f, '') for f in compare if f != 'source')
            action = 'duplicate' if duplicate else 'conflict' if old is not None else 'add'
            output.append(dict(row=row, action=action, existing=old))
            known.setdefault(identity, row)
        except ValueError as exc:
            output.append(dict(row=value, action='error', error=str(exc)))
    return output


def apply_import(kind, document, operations):
    if len(operations) > 5000:
        raise ValueError('最多导入 5000 条')
    rows = [dict(r) for r in document['rows']]
    for operation in operations:
        if not isinstance(operation, dict) or set(operation) != {'action', 'row'}:
            raise ValueError('无效导入操作')
        if operation['action'] == 'skip':
            continue
        if operation['action'] not in ('add', 'replace'):
            raise ValueError('无效导入操作')
        # Preview-only provenance/status cannot override server-owned fields.
        value = {k: v for k, v in operation['row'].items() if k not in ('origin', 'status')}
        row = candidate(kind, value)
        identity = terms.key(kind, row)
        existing = next((i for i, r in enumerate(rows) if terms.key(kind, r) == identity), None)
        if existing is not None:
            if operation['action'] != 'replace':
                raise ValueError('存在冲突，请预览并明确选择替换，或跳过')
            rows[existing] = row
        else:
            rows.append(row)
    result = {**document, 'rows': rows}
    terms.decode(kind, terms.encode(result))
    return result
