"""Constrained citation presentation edits. Never enrich missing metadata."""
from copy import deepcopy
import re

from docx import Document

from .documents import doc_paragraphs, replace_text
from .jobs import object_schema


def apply_citation_edits(path, edits, reference_positions):
    doc = Document(path)
    paragraphs = [p for p in doc_paragraphs(doc) if p.text.strip()]
    issues, applied = [], 0
    for edit in edits:
        i = edit['paragraph'] - 1
        old, new = edit['original'], edit['replacement']
        if not 0 <= i < len(paragraphs) or not old or paragraphs[i].text.count(old) != 1:
            issues.append('Unmatched citation edit was left unchanged.')
            continue
        if i + 1 in reference_positions:
            # Reordering/punctuation is allowed; new bibliographic facts are not.
            vocabulary = set(re.findall(r'\w+', old.casefold())) | {str(n) for n in range(1, len(reference_positions) + 1)} | {'vol', 'no', 'pp', 'p', 'doi', 'in', 'ed', 'eds'}
            if set(re.findall(r'\w+', new.casefold())) - vocabulary:
                issues.append('A reference proposed new metadata; it was left unchanged.')
                continue
        elif not (re.fullmatch(r'(?:\([^()\n]*\d{4}[a-z]?[^()\n]*\)|\[[\d,;\s–—-]+\])', old)
                  and re.fullmatch(r'\[[\d,;\s–—-]+\]', new)):
            issues.append('A proposed body change was not a citation marker; it was left unchanged.')
            continue
        paragraph = paragraphs[i]
        replace_text(paragraph, paragraph.text.replace(old, new, 1))
        # Preserve template run properties while applying explicitly identified title italics.
        italics = [text for text in edit.get('italic', []) if text and text in paragraph.text]
        if italics and i + 1 in reference_positions and not paragraph._p.xpath('.//w:drawing | .//w:object | .//w:fldChar | .//w:hyperlink'):
            text = paragraph.text
            props = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
            paragraph.clear()
            pattern = '(' + '|'.join(re.escape(t) for t in sorted(italics, key=len, reverse=True)) + ')'
            for fragment in re.split(pattern, text):
                run = paragraph.add_run(fragment)
                if props is not None:
                    run._r.insert(0, deepcopy(props))
                run.italic = fragment in italics
        applied += 1
    doc.save(path)
    return applied, issues


async def format_citations(worker, job, path, payload, details):
    from .presets import structure
    data = await worker.blocking(structure, path)
    doc = Document(path)
    paragraphs = [p.text for p in doc_paragraphs(doc) if p.text.strip()]
    references = {row['text'] for row in data['blocks'] if row.get('role') == 'reference'}
    positions = {i for i, text in enumerate(paragraphs, 1) if text in references}
    selected = [{'paragraph': i, 'text': text, 'reference': i in positions} for i, text in enumerate(paragraphs, 1)
                if i in positions or re.search(r'\([^()\n]*\d{4}[^()\n]*\)|\[[\d,;\s–—-]+\]', text)]
    if not selected:
        return {'applied': 0, 'issues': []}
    if not positions:
        return {'applied': 0, 'issues': ['Reference identities cannot be established without a recognized bibliography; citation markers were preserved.']}
    if sum(len(row['text']) for row in selected) > 45000:
        return {'applied': 0, 'issues': ['Citation content exceeds the bounded formatting budget; bibliography presentation requires review.']}
    schema = object_schema({'edits': {'type': 'array', 'items': object_schema({
        'paragraph': {'type': 'integer'}, 'original': {'type': 'string'}, 'replacement': {'type': 'string'},
        'italic': {'type': 'array', 'items': {'type': 'string'}}})},
        'issues': {'type': 'array', 'items': {'type': 'string'}}})
    import json
    worker.phase(job, 'citations', '正在排版引文与参考文献', '仅使用文档已有信息，缺失信息保留待审')
    result = await worker.runner.run(job['project'], job['id'],
        'Format citation presentation for ' + payload['template'] + '. Use numeric bracket citations, preserving unambiguous reference identity. '
        'Keep narrative wording exactly unchanged. For body edits original must be ONLY the citation marker. '
        'For bibliography edits retain all metadata, changing punctuation/order only; identify publication titles for italics. '
        'Do not add, remove or renumber references when identity is ambiguous. No external tools or research. '
        'Flag missing fields rather than filling them. Return issues for unresolved citations.\n' + json.dumps(selected, ensure_ascii=False), schema)
    applied, issues = await worker.blocking(apply_citation_edits, path, result['edits'], positions)
    return {'applied': applied, 'issues': result['issues'] + issues}
