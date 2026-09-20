from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.text.paragraph import Paragraph
from pypdf import PdfReader


def doc_paragraphs(doc):
    # Body order includes paragraphs nested in tables and content controls.
    return [Paragraph(el, doc._body) for el in doc.element.body.xpath(".//w:p")]


def extract(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".docx":
        with ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("DOCX 解压后超过 100 MB")
        blocks = [p.text for p in doc_paragraphs(Document(path)) if p.text.strip()]
    elif suffix == ".pdf":
        pdf = PdfReader(path)
        if pdf.is_encrypted:
            raise ValueError("请先移除 PDF 密码")
        blocks = [b.strip() for page in pdf.pages for b in (page.extract_text() or "").split("\n\n") if b.strip()]
    elif suffix in (".txt", ".md"):
        blocks = [b.strip() for b in Path(path).read_text(encoding="utf-8").splitlines() if b.strip()]
    else:
        raise ValueError("支持 PDF、DOCX、UTF-8 TXT 和 Markdown")
    if not blocks:
        raise ValueError("未提取到文本；扫描 PDF 需要先 OCR")
    if sum(map(len, blocks)) > 2_000_000:
        raise ValueError("文档文本超过 200 万字符，请拆分上传")
    return blocks


def export_docx(source, translations, output):
    if source.suffix.lower() == ".docx":
        doc = Document(source)
        paragraphs = [p for p in doc_paragraphs(doc) if p.text.strip()]
        if len(paragraphs) != len(translations):
            raise ValueError("段落数量不匹配，拒绝导出")
        for p, text in zip(paragraphs, translations):
            # Keep paragraph/table styles, first run style and non-text objects.
            for node in p._p.xpath('.//w:br | .//w:cr | .//w:tab'):
                node.getparent().remove(node)
            text_nodes = p._p.xpath(".//w:t")
            if text_nodes:
                text_nodes[0].text = text
                text_nodes[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                for node in text_nodes[1:]:
                    node.text = ""
            else:
                p.add_run(text)
    else:
        doc = Document()
        for text in translations:
            doc.add_paragraph(text)
    doc.save(output)


def paragraph_structure(source, count):
    """Lightweight reading structure, independent of page layout."""
    fallback = [{"kind": "paragraph"} for _ in range(count)]
    if source.suffix.lower() != ".docx":
        return fallback
    doc = Document(source)
    paragraphs = [p for p in doc_paragraphs(doc) if p.text.strip()]
    if len(paragraphs) != count:
        return fallback
    tables = doc.element.body.xpath(".//w:tbl")
    rows = []
    for paragraph in paragraphs:
        cell = next((el for el in paragraph._p.iterancestors() if el.tag.endswith('}tc')), None)
        if cell is not None:
            table_row = cell.getparent()
            table = next(el for el in table_row.iterancestors() if el.tag.endswith('}tbl'))
            rows.append({"kind": "table", "table": tables.index(table) + 1,
                         "row": table.xpath('.//w:tr').index(table_row) + 1,
                         "column": table_row.xpath('./w:tc').index(cell) + 1})
        else:
            match = re.fullmatch(r'Heading\s*([1-9])', paragraph.style.name or '', re.I)
            rows.append({"kind": "heading", "level": int(match[1])} if match else {"kind": "paragraph"})
    return rows


def revise_docx(source, expected, index, replacement, output):
    """Clone a published version and change exactly one aligned paragraph."""
    doc = Document(source)
    paragraphs = [p for p in doc_paragraphs(doc) if p.text.strip()]
    if [p.text for p in paragraphs] != expected or not 0 <= index < len(paragraphs):
        raise ValueError("译文与对照记录不一致，拒绝修改文档")
    paragraph = paragraphs[index]
    text_nodes = paragraph._p.xpath('.//w:t')
    # Clear textual controls only inside the selected paragraph.
    for node in paragraph._p.xpath('.//w:br | .//w:cr | .//w:tab'):
        node.getparent().remove(node)
    if text_nodes:
        text_nodes[0].text = replacement
        text_nodes[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        for node in text_nodes[1:]:
            node.text = ''
    else:
        paragraph.add_run(replacement)
    doc.save(output)


def legacy_alignment(work, job_id, source, published):
    snapshot = work / 'runs' / job_id / 'alignment.json'
    if snapshot.exists():
        import json
        from .alignment import target_texts
        if snapshot.is_symlink() or not snapshot.resolve().is_relative_to(work.resolve()):
            raise ValueError('无效对照快照路径')
        rows = json.loads(snapshot.read_text())
        if target_texts(rows) != extract(published):
            raise ValueError('译文与对照快照不一致')
        return rows
    inputs = sorted((work / 'runs' / job_id).glob('batch-*-input.json'),
                    key=lambda p: int(p.name.split('-')[1]))
    import json
    originals = []
    for path in inputs:
        if path.is_symlink() or not path.resolve().is_relative_to(work.resolve()):
            raise ValueError('无效原文快照路径')
        originals.extend(json.loads(path.read_text())['source'])
    translated = extract(published)
    if not originals or len(originals) != len(translated):
        raise ValueError('历史译文缺少可靠的段落对应记录，暂不能对照')
    structure = paragraph_structure(source, len(originals))
    return [{**meta, 'original': original, 'translation': translation}
            for meta, original, translation in zip(structure, originals, translated)]


def replace_text(paragraph, text):
    for node in paragraph._p.xpath('.//w:br[not(@w:type) or @w:type="textWrapping"] | .//w:cr | .//w:tab'):
        node.getparent().remove(node)
    nodes = paragraph._p.xpath('.//w:t')
    if nodes:
        nodes[0].text = text
        nodes[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        for node in nodes[1:]:
            node.text = ''
    else:
        paragraph.add_run(text)


def replace_span(paragraphs, texts):
    from copy import deepcopy
    from docx.oxml import OxmlElement
    first = paragraphs[0]
    replace_text(first, texts[0])
    for paragraph in paragraphs[1:]:
        if paragraph._p.getparent() is not first._p.getparent():
            raise ValueError('不能跨容器合并段落')
        paragraph._p.getparent().remove(paragraph._p)
    anchor = first._p
    for text in texts[1:]:
        node = OxmlElement('w:p')
        if first._p.pPr is not None:
            node.append(deepcopy(first._p.pPr))
        p = Paragraph(node, first._parent)
        run = p.add_run(text)
        if first.runs and first.runs[0]._r.rPr is not None:
            run._r.insert(0, deepcopy(first.runs[0]._r.rPr))
        anchor.addnext(node)
        anchor = node


def export_groups(source, blocks, rows, output):
    from .alignment import normalize_groups, target_texts
    normalize_groups([dict(source_ids=r['source_ids'], paragraphs=r['translations'], reason=r['reason']) for r in rows], blocks)
    if source.suffix.lower() == '.docx':
        doc = Document(source)
        paragraphs = [p for p in doc_paragraphs(doc) if p.text.strip()]
        if [p.text for p in paragraphs] != [b['text'] for b in blocks]:
            raise ValueError('原文在翻译期间发生变化，拒绝导出')
        offset = 0
        for row in rows:
            count = len(row['source_ids'])
            replace_span(paragraphs[offset:offset + count], row['translations'])
            offset += count
        doc.save(output)
    else:
        export_docx(source, target_texts(rows), output)
    if extract(output) != target_texts(rows):
        raise ValueError('导出译文与段落组不一致，拒绝发布')


def revise_group_docx(source, rows, index, texts, output):
    from .alignment import target_texts
    doc = Document(source)
    paragraphs = [p for p in doc_paragraphs(doc) if p.text.strip()]
    if [p.text for p in paragraphs] != target_texts(rows):
        raise ValueError('译文与对照记录不一致，拒绝修改文档')
    counts = [len(row.get('translations', [row['translation']])) for row in rows]
    start = sum(counts[:index])
    replace_span(paragraphs[start:start + counts[index]], texts)
    doc.save(output)
    expected = target_texts(rows[:index]) + texts + target_texts(rows[index + 1:])
    if extract(output) != expected:
        raise ValueError('修订导出与段落组不一致，拒绝发布')
