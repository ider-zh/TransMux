"""Opt-in real CLI acceptance test; uses synthetic documents in isolated storage."""
import argparse
import asyncio
import json
from pathlib import Path
import tempfile

from docx import Document

from transmux.alignment import target_texts
from transmux.documents import extract
from transmux.jobs import Worker
from transmux.store import Store


async def main(agent, model):
    root = Path(tempfile.mkdtemp(prefix=f'groups-{agent}-', dir=Path('build').resolve()))
    store = Store(root)
    project = store.create_project('Synthetic paragraph grouping acceptance', agent, model=model)
    pid = project['id']
    work = store.workspace(pid)
    (work / 'style.md').write_text('# English style\nUse clear academic English. Repair obviously broken sentences by merging adjacent body paragraphs. '
                                  'When a paragraph changes from experimental results to a separate discussion of limitations, split those topics into two paragraphs.\n')
    source = work / 'sources' / 'synthetic.docx'
    doc = Document()
    doc.add_heading('实验报告', 1)
    doc.add_paragraph('该模型通过分析大量')
    doc.add_paragraph('训练数据来预测温度。')
    doc.add_paragraph('实验结果显示，预测误差降低了百分之十，计算速度提高了百分之二十。关于局限性，本实验只使用了一个地区的数据，因此还需要跨地区验证。')
    doc.add_paragraph('样本数量为一百。', style='List Bullet')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '温度：二十摄氏度。'
    doc.add_heading('结论', 1)
    doc.add_paragraph('该研究为后续的')
    doc.add_paragraph('跨地区验证提供了依据。')
    doc.save(source)
    original = source.read_bytes()
    Path(str(source) + '.json').write_text(json.dumps(extract(source), ensure_ascii=False))
    fid = store.add_file(pid, source.name, 'source', source)
    worker = Worker(store)

    async def execute(kind, payload):
        task = store.enqueue(pid, kind, {**payload, 'use_rag': False, 'max_review_rounds': 3})
        print(json.dumps({'agent': agent, 'root': str(root), 'job': task['id'], 'kind': kind}), flush=True)
        await worker.execute(task)
        finished = store.rows('SELECT state,result FROM jobs WHERE id=?', (task['id'],))[0]
        print(json.dumps(finished, ensure_ascii=False), flush=True)
        assert finished['state'] == 'succeeded', finished
        return json.loads(finished['result'])['file_id']

    output = await execute('translate', {'file_id': fid})
    rows = store.comparison(pid, output)['paragraphs']
    merged = next(i for i, r in enumerate(rows) if r['source_ids'] == ['p000002', 'p000003'])
    assert len(rows[merged]['translations']) == 1
    assert any(r['source_ids'] == ['p000004'] and len(r['translations']) == 2 for r in rows)
    assert all(len(r['translations']) == 1 for r in rows if r['protected'])
    assert any(r['source_ids'] == ['p000008', 'p000009'] and len(r['translations']) == 1 for r in rows)
    inputs = sorted((work / 'runs').glob('*/batch-*-input.json'))
    requests = [json.loads(p.read_text()) for p in inputs]
    assert sorted(len(r['source']) for r in requests) == [3, 6]
    assert all(r['section_path'] and not r['continuation'] for r in requests)
    path = store.download_path(pid, store.file(pid, output))
    assert extract(path) == target_texts(rows)
    before = path.read_bytes()
    revised = await execute('revise', {'file_id': output, 'paragraph': merged + 1,
                                     'message': 'Restore two target paragraphs corresponding to the two original source paragraphs. Keep all meaning, use two natural English sentences even though the original break was awkward.'})
    revised_rows = store.comparison(pid, revised)['paragraphs']
    assert len(revised_rows[merged]['translations']) == 2
    assert revised_rows[:merged] == rows[:merged] and revised_rows[merged + 1:] == rows[merged + 1:]
    assert source.read_bytes() == original and path.read_bytes() == before
    report = dict(agent=agent, model=model, state='passed', root=str(root),
                  translation=str(path), revision=str(store.download_path(pid, store.file(pid, revised))),
                  source_paragraphs=len(extract(source)), groups=len(rows), target_paragraphs=len(target_texts(rows)))
    (root / 'acceptance.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report, ensure_ascii=False), flush=True)
    store.db.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('agent', choices=['codex', 'codebuddy'])
    parser.add_argument('--model', required=True)
    args = parser.parse_args()
    asyncio.run(main(args.agent, args.model))
